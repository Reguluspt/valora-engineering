"""PostgreSQL-only concurrency and lock-order proof for PR-01 preliminary result generation."""
from __future__ import annotations

import os
import threading
import uuid

import pytest
from fastapi import HTTPException
from openpyxl import Workbook
from sqlalchemy import create_engine, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, sessionmaker

from app.modules.excel_import.infrastructure.object_storage import (
    FakeObjectStorage,
    set_object_storage_override,
)
from app.modules.excel_import.models import (
    ColumnMappingDecision,
    ColumnMappingDecisionKind,
    ColumnMappingDecisionOutcome,
    ColumnMappingMemoryScope,
    ColumnMappingProfileUsage,
    ColumnMappingProposalSourceKind,
    ImportSourceArtifact,
    ImportSourceArtifactState,
    WorkbookStructureDisposition,
    WorkbookStructureSnapshot,
)
from app.modules.project_master_data.application.official_intake_service import (
    commit_project_official_intake,
)
from app.modules.project_master_data.application.preliminary_result_service import (
    generate_preliminary_result_artifact,
)
import app.modules.project_master_data.application.preliminary_result_service as result_service
from app.modules.project_master_data.models import (
    AuditEvent,
    Customer,
    CustomerStatus,
    ImportBatchStatus,
    OrganizationProfile,
    OrganizationStatus,
    PreliminaryAnalysisSnapshot,
    PreliminaryResultArtifact,
    Project,
    ProjectAssetImportBatch,
    ProjectOfficialIntakeCommit,
    ProjectWorkflowStatus,
    Role,
    User,
    UserRole,
    UserStatus,
)


RESULT_PERMISSION = "project:preliminary_result:generate"
OFFICIAL_INTAKE_PERMISSION = "project:official_intake:commit"


def _postgres_engine_or_skip():
    url = os.getenv("TEST_DATABASE_URL")
    if not url or not url.startswith("postgres"):
        if os.getenv("CI") == "true":
            pytest.fail("CI=true requires PostgreSQL TEST_DATABASE_URL for PR-01 preliminary result")
        pytest.skip("PR-01 preliminary result concurrency proof requires PostgreSQL TEST_DATABASE_URL")
    engine = create_engine(url, connect_args={"connect_timeout": 5}, pool_pre_ping=True)
    with engine.connect() as connection:
        exists = connection.execute(
            text("SELECT to_regclass('preliminary_result_artifacts')")
        ).scalar_one()
    if exists is None:
        engine.dispose()
        if os.getenv("CI") == "true":
            pytest.fail("CI PostgreSQL is not migrated to the PR-01 preliminary result head")
        pytest.skip("PostgreSQL is not migrated to the PR-01 preliminary result head")
    return engine


def _make_xlsx_bytes(*, rows: list[tuple[str, float]]) -> bytes:
    import io

    wb = Workbook()
    ws = wb.active
    ws.title = "Assets"
    ws.cell(row=1, column=1, value="Tên tài sản")
    ws.cell(row=2, column=1, value="Tên tài sản")
    ws.cell(row=1, column=2, value="Số lượng")
    ws.cell(row=2, column=2, value="Số lượng")
    ws.cell(row=1, column=3, value="Ghi chú")
    ws.cell(row=2, column=3, value="Ghi chú")
    for index, (name, qty) in enumerate(rows, start=3):
        ws.cell(row=index, column=1, value=name)
        ws.cell(row=index, column=2, value=qty)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()


def _mapping_snapshot(*, artifact_id, generation, source_checksum, structure_id, max_row: int):
    return {
        "contract_version": "s13-pr-004-v1",
        "source": {
            "source_artifact_id": str(artifact_id),
            "generation": generation,
            "checksum_sha256": source_checksum,
        },
        "structure": {
            "structure_snapshot_id": str(structure_id),
            "snapshot_version": 1,
            "rule_version": "s13-pr-003-v3",
            "analysis_digest_sha256": "b" * 64,
        },
        "template_fingerprint_sha256": "c" * 64,
        "candidate": {
            "candidate_index": 0,
            "sheet_name": "Assets",
            "header_start_row": 1,
            "header_end_row": 2,
            "data_start_row": 3,
            "min_row": 1,
            "max_row": max_row,
            "min_column": 1,
            "max_column": 3,
        },
        "fields": [
            {
                "source_column_index": 1,
                "source_column_letter": "A",
                "original_header": "Tên tài sản",
                "semantic_role": "raw_asset_name",
            },
            {
                "source_column_index": 2,
                "source_column_letter": "B",
                "original_header": "Số lượng",
                "semantic_role": "quantity",
            },
            {
                "source_column_index": 3,
                "source_column_letter": "C",
                "original_header": "Ghi chú",
                "semantic_role": "ignore",
            },
        ],
    }


def _seed(setup: Session, *, suffix: str):
    rows = [("Máy bơm 10HP", 2.0)]
    source_data = _make_xlsx_bytes(rows=rows)
    import hashlib

    source_checksum = hashlib.sha256(source_data).hexdigest()

    org = OrganizationProfile(
        legal_name=f"PG Preliminary Result Org {suffix}",
        organization_slug=f"pg-preliminary-result-{suffix}-{uuid.uuid4().hex[:8]}",
        status=OrganizationStatus.ACTIVE,
    )
    setup.add(org)
    setup.flush()
    actor = User(
        organization_id=org.id,
        email=f"pg-preliminary-result-{suffix}-{uuid.uuid4().hex[:8]}@example.com",
        full_name="PG Preliminary Result Human",
        status=UserStatus.ACTIVE,
    )
    setup.add(actor)
    setup.flush()
    role = Role(
        code=f"pg-preliminary-result-{suffix}-{uuid.uuid4().hex[:8]}",
        display_name="PG Preliminary Result Fixture",
        permissions=[RESULT_PERMISSION, OFFICIAL_INTAKE_PERMISSION],
    )
    setup.add(role)
    setup.flush()
    setup.add(UserRole(user_id=actor.id, role_id=role.id, is_active=True))
    customer = Customer(
        organization_id=org.id,
        legal_name=f"PG Preliminary Result Customer {suffix}",
        status=CustomerStatus.ACTIVE,
        created_by=actor.id,
    )
    setup.add(customer)
    setup.flush()
    project = Project(
        organization_id=org.id,
        customer_id=customer.id,
        code=f"PG-PR-{suffix}-{uuid.uuid4().hex[:6]}",
        name=f"PG Preliminary Result Project {suffix}",
        status=ProjectWorkflowStatus.DRAFT,
        created_by=actor.id,
    )
    setup.add(project)
    setup.flush()
    batch = ProjectAssetImportBatch(
        organization_id=org.id,
        project_id=project.id,
        source_filename="danh-muc-v1.xlsx",
        status=ImportBatchStatus.PARSED.value,
        total_rows=1,
        valid_rows=1,
        created_by_user_id=actor.id,
    )
    setup.add(batch)
    setup.flush()
    artifact = ImportSourceArtifact(
        organization_id=org.id,
        project_id=project.id,
        import_batch_id=batch.id,
        generation=1,
        original_filename="danh-muc-v1.xlsx",
        detected_format="xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        file_size_bytes=len(source_data),
        checksum_sha256=source_checksum,
        storage_object_key=f"import-sources/{project.id}/v1.xlsx",
        state=ImportSourceArtifactState.AVAILABLE.value,
        created_by_user_id=actor.id,
    )
    setup.add(artifact)
    setup.flush()
    batch.current_source_artifact_id = artifact.id
    setup.flush()
    structure = WorkbookStructureSnapshot(
        organization_id=org.id,
        project_id=project.id,
        import_batch_id=batch.id,
        source_artifact_id=artifact.id,
        snapshot_version=1,
        source_checksum_sha256=source_checksum,
        rule_version="s13-pr-003-v3",
        adapter_name="excel",
        adapter_version="1.0",
        disposition=WorkbookStructureDisposition.PROPOSED.value,
        candidate_count=1,
        structure_payload={"sheet": "Assets"},
        analysis_digest_sha256="b" * 64,
        created_by_user_id=actor.id,
    )
    setup.add(structure)
    setup.flush()
    proposal_decision = ColumnMappingDecision(
        organization_id=org.id,
        customer_id=customer.id,
        project_id=project.id,
        import_batch_id=batch.id,
        source_artifact_id=artifact.id,
        structure_snapshot_id=structure.id,
        decision_kind=ColumnMappingDecisionKind.PROPOSAL.value,
        outcome=ColumnMappingDecisionOutcome.PROPOSED.value,
        memory_scope=ColumnMappingMemoryScope.NONE.value,
        actor_user_id=actor.id,
        command_id=uuid.uuid4(),
        proposal_source_kind=ColumnMappingProposalSourceKind.HUMAN.value,
        proposal_source_version="1.0",
        mapping_contract_version="v1",
        template_fingerprint_sha256="c" * 64,
        mapping_snapshot={"roles": []},
        mapping_digest_sha256="d" * 64,
        before_summary={},
        after_summary={},
    )
    setup.add(proposal_decision)
    setup.flush()
    decision = ColumnMappingDecision(
        organization_id=org.id,
        customer_id=customer.id,
        project_id=project.id,
        import_batch_id=batch.id,
        source_artifact_id=artifact.id,
        structure_snapshot_id=structure.id,
        decision_kind=ColumnMappingDecisionKind.CONFIRMATION.value,
        outcome=ColumnMappingDecisionOutcome.ACCEPTED.value,
        memory_scope=ColumnMappingMemoryScope.NONE.value,
        proposal_decision_id=proposal_decision.id,
        actor_user_id=actor.id,
        command_id=uuid.uuid4(),
        proposal_source_kind=ColumnMappingProposalSourceKind.HUMAN.value,
        proposal_source_version="1.0",
        mapping_contract_version="v1",
        template_fingerprint_sha256="c" * 64,
        mapping_snapshot=_mapping_snapshot(
            artifact_id=artifact.id,
            generation=1,
            source_checksum=source_checksum,
            structure_id=structure.id,
            max_row=2 + len(rows),
        ),
        mapping_digest_sha256="d" * 64,
        before_summary={},
        after_summary={},
    )
    setup.add(decision)
    setup.flush()
    usage = ColumnMappingProfileUsage(
        organization_id=org.id,
        customer_id=customer.id,
        project_id=project.id,
        import_batch_id=batch.id,
        source_artifact_id=artifact.id,
        structure_snapshot_id=structure.id,
        confirmation_decision_id=decision.id,
        command_id=uuid.uuid4(),
        materialization_contract_version="v1",
        mapping_contract_version="v1",
        template_fingerprint_sha256="c" * 64,
        mapping_snapshot=decision.mapping_snapshot,
        mapping_digest_sha256="d" * 64,
        source_checksum_sha256=source_checksum,
        structure_digest_sha256="b" * 64,
        materialized_asset_row_count=1,
        created_by_user_id=actor.id,
    )
    setup.add(usage)
    setup.flush()

    line_manifest = [
        {
            "identity": "Máy bơm 10HP",
            "accepted_price_basis": "internet_survey",
            "confirmed_reference_price": 1_000_000.0,
            "transport_percentage": 5.0,
            "proposed_unit_price": 1_050_000.0,
            "human_line_confirmed": True,
            "has_unresolved_blocking_line": False,
            "source_row_number": 3,
            "quantity": 2.0,
        }
    ]
    snapshot = PreliminaryAnalysisSnapshot(
        organization_id=org.id,
        customer_id=customer.id,
        project_id=project.id,
        version=1,
        import_batch_id=batch.id,
        source_artifact_id=artifact.id,
        structure_snapshot_id=structure.id,
        mapping_decision_id=decision.id,
        mapping_profile_usage_id=usage.id,
        source_artifact_generation=artifact.generation,
        mapping_decision_digest_sha256=decision.mapping_digest_sha256,
        profile_usage_mapping_digest_sha256=usage.mapping_digest_sha256,
        line_manifest=line_manifest,
        line_manifest_digest_sha256=result_service._sha256_hex(
            result_service._canonical_json(line_manifest)
        ),
        finalized_by_user_id=actor.id,
        idempotency_key=f"pg-preliminary-analysis-{suffix}",
        request_digest_sha256="e" * 64,
    )
    setup.add(snapshot)
    setup.commit()

    fake_storage = FakeObjectStorage()
    fake_storage._objects[artifact.storage_object_key] = source_data
    fake_storage._content_types[artifact.storage_object_key] = artifact.content_type

    snapshot_digest = result_service._snapshot_canonical_digest(snapshot)
    request_digest = result_service._request_digest(
        actor_id=actor.id,
        project_id=project.id,
        snapshot_id=snapshot.id,
        snapshot_digest=snapshot_digest,
        expected_project_version=project.row_version,
    )
    expected_artifact_id = result_service._artifact_id(
        org_id=org.id,
        normalized_key="pg-preliminary-result-key",
        request_digest=request_digest,
    )

    return {
        "org_id": org.id,
        "customer_id": customer.id,
        "actor_id": actor.id,
        "project_id": project.id,
        "snapshot_id": snapshot.id,
        "usage_id": usage.id,
        "expected_artifact_id": expected_artifact_id,
        "fake_storage": fake_storage,
    }


def _run_pair(SessionLocal, *, work: list[tuple]):
    barrier = threading.Barrier(2, timeout=30)
    results: list[uuid.UUID] = []
    errors: list[BaseException] = []

    def worker(callable, *args):
        db: Session = SessionLocal()
        try:
            barrier.wait(timeout=30)
            result = callable(db, *args)
            results.append(result.id if hasattr(result, "id") else result)
        except BaseException as exc:
            errors.append(exc)
        finally:
            db.close()

    threads = [
        threading.Thread(target=worker, args=item, name=f"pr01-result-worker-{index}", daemon=True)
        for index, item in enumerate(work)
    ]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join(timeout=60)
    assert all(not thread.is_alive() for thread in threads)
    return results, errors


def test_postgresql_snapshot_usage_fk_rejects_mismatched_usage_identity() -> None:
    engine = _postgres_engine_or_skip()
    SessionLocal = sessionmaker(bind=engine, expire_on_commit=False)
    setup: Session = SessionLocal()
    try:
        first = _seed(setup, suffix=f"pg-usage-fk-first-{uuid.uuid4().hex[:8]}")
        second = _seed(setup, suffix=f"pg-usage-fk-second-{uuid.uuid4().hex[:8]}")
        snapshot = setup.get(PreliminaryAnalysisSnapshot, first["snapshot_id"])
        assert snapshot is not None
        snapshot.mapping_profile_usage_id = second["usage_id"]

        with pytest.raises(IntegrityError):
            setup.commit()
        setup.rollback()
    finally:
        setup.close()
        engine.dispose()


def test_postgresql_concurrent_same_request_replays_with_one_artifact() -> None:
    engine = _postgres_engine_or_skip()
    SessionLocal = sessionmaker(bind=engine, expire_on_commit=False)
    setup: Session = SessionLocal()
    ids = None
    try:
        ids = _seed(setup, suffix=f"pg-same-{uuid.uuid4().hex[:8]}")
        set_object_storage_override(ids["fake_storage"])
        setup.close()

        def generate(db: Session, actor_id: uuid.UUID):
            actor = db.get(User, actor_id)
            return generate_preliminary_result_artifact(
                db,
                actor=actor,
                org_id=ids["org_id"],
                project_id=ids["project_id"],
                preliminary_analysis_snapshot_id=ids["snapshot_id"],
                expected_project_version=1,
                idempotency_key="pg-preliminary-result-key",
                confirmed=True,
                correlation_id="pg-same",
            )

        results, errors = _run_pair(
            SessionLocal,
            work=[
                (generate, ids["actor_id"]),
                (generate, ids["actor_id"]),
            ],
        )

        assert errors == []
        assert len(results) == 2
        assert results[0] == results[1]
        verify: Session = SessionLocal()
        try:
            assert verify.query(PreliminaryResultArtifact).filter_by(
                organization_id=ids["org_id"], project_id=ids["project_id"]
            ).count() == 1
            assert verify.query(AuditEvent).filter_by(
                organization_id=ids["org_id"], event_name="PreliminaryResultArtifactGenerated"
            ).count() == 1
        finally:
            verify.close()
    finally:
        setup.close()
        engine.dispose()


def test_postgresql_concurrent_same_key_different_digest_is_typed_reuse() -> None:
    engine = _postgres_engine_or_skip()
    SessionLocal = sessionmaker(bind=engine, expire_on_commit=False)
    setup: Session = SessionLocal()
    ids = None
    try:
        ids = _seed(setup, suffix=f"pg-reuse-{uuid.uuid4().hex[:8]}")
        second_actor = User(
            organization_id=ids["org_id"],
            email=f"pg-reuse-peer-{uuid.uuid4().hex[:8]}@example.com",
            full_name="PG Preliminary Result Peer",
        )
        setup.add(second_actor)
        setup.flush()
        role = setup.query(Role).filter(Role.code.like("pg-preliminary-result-%")).first()
        setup.add(UserRole(user_id=second_actor.id, role_id=role.id, is_active=True))
        setup.commit()
        second_actor_id = second_actor.id
        set_object_storage_override(ids["fake_storage"])
        setup.close()

        def generate(db: Session, actor_id: uuid.UUID):
            actor = db.get(User, actor_id)
            return generate_preliminary_result_artifact(
                db,
                actor=actor,
                org_id=ids["org_id"],
                project_id=ids["project_id"],
                preliminary_analysis_snapshot_id=ids["snapshot_id"],
                expected_project_version=1,
                idempotency_key="pg-preliminary-result-key",
                confirmed=True,
                correlation_id="pg-reuse",
            )

        results, errors = _run_pair(
            SessionLocal,
            work=[
                (generate, ids["actor_id"]),
                (generate, second_actor_id),
            ],
        )

        assert len(results) == 1
        assert len(errors) == 1
        assert isinstance(errors[0], HTTPException)
        assert errors[0].status_code == 409
        assert errors[0].detail["error_code"] == "preliminary_result_idempotency_key_reused"
    finally:
        setup.close()
        engine.dispose()


def test_postgresql_concurrent_different_keys_have_one_typed_already_generated() -> None:
    engine = _postgres_engine_or_skip()
    SessionLocal = sessionmaker(bind=engine, expire_on_commit=False)
    setup: Session = SessionLocal()
    ids = None
    try:
        ids = _seed(setup, suffix=f"pg-project-{uuid.uuid4().hex[:8]}")
        set_object_storage_override(ids["fake_storage"])
        setup.close()

        def generate(db: Session, actor_id: uuid.UUID, key: str):
            actor = db.get(User, actor_id)
            return generate_preliminary_result_artifact(
                db,
                actor=actor,
                org_id=ids["org_id"],
                project_id=ids["project_id"],
                preliminary_analysis_snapshot_id=ids["snapshot_id"],
                expected_project_version=1,
                idempotency_key=key,
                confirmed=True,
                correlation_id="pg-project",
            )

        results, errors = _run_pair(
            SessionLocal,
            work=[
                (generate, ids["actor_id"], "pg-project-key-a"),
                (generate, ids["actor_id"], "pg-project-key-b"),
            ],
        )

        assert len(results) == 1
        assert len(errors) == 1
        assert isinstance(errors[0], HTTPException)
        assert errors[0].status_code == 409
        assert errors[0].detail["error_code"] == "preliminary_result_already_generated"
    finally:
        setup.close()
        engine.dispose()


def test_postgresql_concurrent_generate_vs_official_intake_interleave() -> None:
    engine = _postgres_engine_or_skip()
    SessionLocal = sessionmaker(bind=engine, expire_on_commit=False)
    setup: Session = SessionLocal()
    ids = None
    try:
        ids = _seed(setup, suffix=f"pg-interleave-{uuid.uuid4().hex[:8]}")
        set_object_storage_override(ids["fake_storage"])
        setup.close()

        def generate(db: Session, actor_id: uuid.UUID):
            actor = db.get(User, actor_id)
            return generate_preliminary_result_artifact(
                db,
                actor=actor,
                org_id=ids["org_id"],
                project_id=ids["project_id"],
                preliminary_analysis_snapshot_id=ids["snapshot_id"],
                expected_project_version=1,
                idempotency_key="pg-preliminary-result-key",
                confirmed=True,
                correlation_id="pg-interleave-gen",
            )

        def commit(db: Session, actor_id: uuid.UUID):
            actor = db.get(User, actor_id)
            return commit_project_official_intake(
                db,
                actor=actor,
                org_id=ids["org_id"],
                project_id=ids["project_id"],
                preliminary_result_artifact_id=ids["expected_artifact_id"],
                expected_project_version=1,
                expected_preliminary_result_version=1,
                idempotency_key="pg-official-intake-key",
                confirmed=True,
                correlation_id="pg-interleave-commit",
            )

        results, errors = _run_pair(
            SessionLocal,
            work=[
                (generate, ids["actor_id"]),
                (commit, ids["actor_id"]),
            ],
        )

        # One succeeds, the other either succeeds or fails with not_found/race.
        assert len(results) >= 1
        assert len(errors) <= 1
        if errors:
            assert isinstance(errors[0], HTTPException)
            assert errors[0].status_code in (404, 409)
        verify: Session = SessionLocal()
        try:
            assert verify.query(PreliminaryResultArtifact).filter_by(
                organization_id=ids["org_id"], project_id=ids["project_id"]
            ).count() == 1
            # Official intake commits exactly once when it wins the race.
            commit_count = verify.query(ProjectOfficialIntakeCommit).filter_by(
                organization_id=ids["org_id"], project_id=ids["project_id"]
            ).count()
            assert commit_count in (0, 1)
        finally:
            verify.close()
    finally:
        setup.close()
        engine.dispose()


@pytest.mark.parametrize(
    ("idempotency_key", "request_digest_sha256"),
    [
        ("unpaired-key", None),
        (None, "a" * 64),
    ],
)
def test_postgresql_unpaired_idempotency_key_and_digest_is_rejected_by_orm(
    idempotency_key: str | None,
    request_digest_sha256: str | None,
) -> None:
    """Pairing CHECK rejects key without digest and digest without key on migrated PostgreSQL."""
    engine = _postgres_engine_or_skip()
    SessionLocal = sessionmaker(bind=engine, expire_on_commit=False)
    setup: Session = SessionLocal()
    try:
        ids = _seed(setup, suffix=f"pg-unpaired-{uuid.uuid4().hex[:8]}")
        artifact = PreliminaryResultArtifact(
            organization_id=ids["org_id"],
            customer_id=ids["customer_id"],
            project_id=ids["project_id"],
            version=1,
            original_filename="unpaired.xlsx",
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            file_size_bytes=1,
            content_checksum_sha256="a" * 64,
            storage_object_key=f"unpaired/{ids['project_id']}.xlsx",
            source_snapshot_sha256="b" * 64,
            lineage_manifest={"contract": "unpaired"},
            created_by_user_id=ids["actor_id"],
            idempotency_key=idempotency_key,
            request_digest_sha256=request_digest_sha256,
        )
        setup.add(artifact)
        with pytest.raises(IntegrityError):
            setup.commit()
        setup.rollback()
    finally:
        setup.close()
        engine.dispose()
