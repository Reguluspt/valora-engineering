"""Focused proof for the PR-01 PreliminaryResultArtifact generation command."""
from __future__ import annotations

import hashlib
import io
import uuid
from datetime import datetime, timezone
from decimal import Decimal

import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.db import Base
import app.modules.excel_import.models  # noqa: F401
import app.modules.project_master_data.application.preliminary_result_service as service
from app.modules.excel_import.infrastructure.object_storage import (
    FakeObjectStorage,
    set_object_storage_override,
)
from app.modules.excel_import.models import (
    ColumnMappingDecision,
    ColumnMappingDecisionKind,
    ColumnMappingDecisionOutcome,
    ColumnMappingMemoryScope,
    ColumnMappingProfileUsage,
    ColumnMappingProposalSourceKind,
    ImportSourceArtifact,
    ImportSourceArtifactState,
    WorkbookStructureDisposition,
    WorkbookStructureSnapshot,
)
from app.modules.project_master_data.application.preliminary_result_service import (
    generate_preliminary_result_artifact,
)
from app.modules.project_master_data.models import (
    AuditEvent,
    Customer,
    CustomerStatus,
    ImportBatchStatus,
    OrganizationProfile,
    OrganizationStatus,
    PreliminaryAnalysisSnapshot,
    PreliminaryResultArtifact,
    Project,
    ProjectAssetImportBatch,
    ProjectWorkflowStatus,
    Role,
    User,
    UserRole,
    UserStatus,
)


PRELIMINARY_RESULT_PERMISSION = "project:preliminary_result:generate"


@pytest.fixture
def result_db() -> Session:
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    session = Session(bind=engine)
    try:
        yield session
    finally:
        session.close()
        Base.metadata.drop_all(engine)


@pytest.fixture(autouse=True)
def _reset_storage_override():
    set_object_storage_override(None)
    yield
    set_object_storage_override(None)


def _make_xlsx_bytes(*, rows: list[tuple[str, float]]) -> bytes:
    """Build a minimal xlsx with two header rows and data starting at row 3."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Assets"
    ws.cell(row=1, column=1, value="Tên tài sản")
    ws.cell(row=2, column=1, value="Tên tài sản")
    ws.cell(row=1, column=2, value="Số lượng")
    ws.cell(row=2, column=2, value="Số lượng")
    ws.cell(row=1, column=3, value="Ghi chú")
    ws.cell(row=2, column=3, value="Ghi chú")
    for index, (name, qty) in enumerate(rows, start=3):
        ws.cell(row=index, column=1, value=name)
        ws.cell(row=index, column=2, value=qty)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _mapping_snapshot(
    *,
    artifact_id: uuid.UUID,
    generation: int,
    source_checksum: str,
    structure_id: uuid.UUID,
    structure_digest: str,
    max_column: int = 3,
    max_row: int = 5,
) -> dict:
    return {
        "contract_version": "s13-pr-004-v1",
        "source": {
            "source_artifact_id": str(artifact_id),
            "generation": generation,
            "checksum_sha256": source_checksum,
        },
        "structure": {
            "structure_snapshot_id": str(structure_id),
            "snapshot_version": 1,
            "rule_version": "s13-pr-003-v3",
            "analysis_digest_sha256": structure_digest,
        },
        "template_fingerprint_sha256": "c" * 64,
        "candidate": {
            "candidate_index": 0,
            "sheet_name": "Assets",
            "header_start_row": 1,
            "header_end_row": 2,
            "data_start_row": 3,
            "min_row": 1,
            "max_row": max_row,
            "min_column": 1,
            "max_column": max_column,
        },
        "fields": [
            {
                "source_column_index": 1,
                "source_column_letter": "A",
                "original_header": "Tên tài sản",
                "semantic_role": "raw_asset_name",
            },
            {
                "source_column_index": 2,
                "source_column_letter": "B",
                "original_header": "Số lượng",
                "semantic_role": "quantity",
            },
            {
                "source_column_index": 3,
                "source_column_letter": "C",
                "original_header": "Ghi chú",
                "semantic_role": "ignore",
            },
        ],
    }


def _line(
    *,
    identity: str = "Máy bơm nước 10HP",
    source_row_number: int = 3,
    quantity: float = 2.0,
    proposed_unit_price: float = 1_050_000.0,
) -> dict:
    return {
        "identity": identity,
        "accepted_price_basis": "internet_survey",
        "confirmed_reference_price": 1_000_000.0,
        "transport_percentage": 5.0,
        "proposed_unit_price": proposed_unit_price,
        "human_line_confirmed": True,
        "has_unresolved_blocking_line": False,
        "source_row_number": source_row_number,
        "quantity": quantity,
    }


def _seed(
    result_db: Session,
    *,
    suffix: str = "a",
    grant_permission: bool = True,
    artifact_state: ImportSourceArtifactState = ImportSourceArtifactState.AVAILABLE,
    rows: list[tuple[str, float]] | None = None,
) -> dict:
    rows = rows if rows is not None else [("Máy bơm nước 10HP", 2.0)]
    source_data = _make_xlsx_bytes(rows=rows)
    source_checksum = _sha256(source_data)

    org = OrganizationProfile(
        legal_name=f"Preliminary Result Org {suffix}",
        organization_slug=f"preliminary-result-{suffix}-{uuid.uuid4().hex[:8]}",
        status=OrganizationStatus.ACTIVE,
    )
    result_db.add(org)
    result_db.flush()
    actor = User(
        organization_id=org.id,
        email=f"preliminary-result-{suffix}-{uuid.uuid4().hex[:8]}@example.com",
        full_name="Preliminary Result Human",
        status=UserStatus.ACTIVE,
    )
    result_db.add(actor)
    result_db.flush()
    role = Role(
        code=f"preliminary-result-{suffix}-{uuid.uuid4().hex[:8]}",
        display_name="Preliminary Result Fixture",
        permissions=[PRELIMINARY_RESULT_PERMISSION] if grant_permission else [],
    )
    result_db.add(role)
    result_db.flush()
    result_db.add(UserRole(user_id=actor.id, role_id=role.id, is_active=True))
    customer = Customer(
        organization_id=org.id,
        legal_name=f"Preliminary Result Customer {suffix}",
        status=CustomerStatus.ACTIVE,
        created_by=actor.id,
    )
    result_db.add(customer)
    result_db.flush()
    project = Project(
        organization_id=org.id,
        customer_id=customer.id,
        code=f"PR-{suffix}-{uuid.uuid4().hex[:6]}",
        name=f"Preliminary Result Project {suffix}",
        status=ProjectWorkflowStatus.DRAFT,
        created_by=actor.id,
    )
    result_db.add(project)
    result_db.flush()
    batch = ProjectAssetImportBatch(
        organization_id=org.id,
        project_id=project.id,
        source_filename="danh-muc-v1.xlsx",
        status=ImportBatchStatus.PARSED.value,
        total_rows=len(rows),
        valid_rows=len(rows),
        created_by_user_id=actor.id,
    )
    result_db.add(batch)
    result_db.flush()
    artifact = ImportSourceArtifact(
        organization_id=org.id,
        project_id=project.id,
        import_batch_id=batch.id,
        generation=1,
        original_filename="danh-muc-v1.xlsx",
        detected_format="xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        file_size_bytes=len(source_data),
        checksum_sha256=source_checksum,
        storage_object_key=f"import-sources/{project.id}/v1.xlsx",
        state=artifact_state.value,
        created_by_user_id=actor.id,
    )
    result_db.add(artifact)
    result_db.flush()
    batch.current_source_artifact_id = artifact.id
    result_db.flush()
    structure = WorkbookStructureSnapshot(
        organization_id=org.id,
        project_id=project.id,
        import_batch_id=batch.id,
        source_artifact_id=artifact.id,
        snapshot_version=1,
        source_checksum_sha256=source_checksum,
        rule_version="s13-pr-003-v3",
        adapter_name="excel",
        adapter_version="1.0",
        disposition=WorkbookStructureDisposition.PROPOSED.value,
        candidate_count=1,
        structure_payload={"sheet": "Assets"},
        analysis_digest_sha256="b" * 64,
        created_by_user_id=actor.id,
    )
    result_db.add(structure)
    result_db.flush()
    proposal_decision = ColumnMappingDecision(
        organization_id=org.id,
        customer_id=customer.id,
        project_id=project.id,
        import_batch_id=batch.id,
        source_artifact_id=artifact.id,
        structure_snapshot_id=structure.id,
        decision_kind=ColumnMappingDecisionKind.PROPOSAL.value,
        outcome=ColumnMappingDecisionOutcome.PROPOSED.value,
        memory_scope=ColumnMappingMemoryScope.NONE.value,
        actor_user_id=actor.id,
        command_id=uuid.uuid4(),
        proposal_source_kind=ColumnMappingProposalSourceKind.HUMAN.value,
        proposal_source_version="1.0",
        mapping_contract_version="v1",
        template_fingerprint_sha256="c" * 64,
        mapping_snapshot={"roles": []},
        mapping_digest_sha256="d" * 64,
        before_summary={},
        after_summary={},
    )
    result_db.add(proposal_decision)
    result_db.flush()
    decision = ColumnMappingDecision(
        organization_id=org.id,
        customer_id=customer.id,
        project_id=project.id,
        import_batch_id=batch.id,
        source_artifact_id=artifact.id,
        structure_snapshot_id=structure.id,
        decision_kind=ColumnMappingDecisionKind.CONFIRMATION.value,
        outcome=ColumnMappingDecisionOutcome.ACCEPTED.value,
        memory_scope=ColumnMappingMemoryScope.NONE.value,
        proposal_decision_id=proposal_decision.id,
        actor_user_id=actor.id,
        command_id=uuid.uuid4(),
        proposal_source_kind=ColumnMappingProposalSourceKind.HUMAN.value,
        proposal_source_version="1.0",
        mapping_contract_version="v1",
        template_fingerprint_sha256="c" * 64,
        mapping_snapshot=_mapping_snapshot(
            artifact_id=artifact.id,
            generation=1,
            source_checksum=source_checksum,
            structure_id=structure.id,
            structure_digest="b" * 64,
            max_row=2 + len(rows),
        ),
        mapping_digest_sha256="d" * 64,
        before_summary={},
        after_summary={},
    )
    result_db.add(decision)
    result_db.flush()
    usage = ColumnMappingProfileUsage(
        organization_id=org.id,
        customer_id=customer.id,
        project_id=project.id,
        import_batch_id=batch.id,
        source_artifact_id=artifact.id,
        structure_snapshot_id=structure.id,
        confirmation_decision_id=decision.id,
        command_id=uuid.uuid4(),
        materialization_contract_version="v1",
        mapping_contract_version="v1",
        template_fingerprint_sha256="c" * 64,
        mapping_snapshot=decision.mapping_snapshot,
        mapping_digest_sha256="d" * 64,
        source_checksum_sha256=source_checksum,
        structure_digest_sha256="b" * 64,
        materialized_asset_row_count=len(rows),
        created_by_user_id=actor.id,
    )
    result_db.add(usage)
    result_db.flush()

    line_manifest = [_line(source_row_number=2 + index, quantity=qty) for index, (_, qty) in enumerate(rows, start=1)]
    snapshot = PreliminaryAnalysisSnapshot(
        organization_id=org.id,
        customer_id=customer.id,
        project_id=project.id,
        version=1,
        import_batch_id=batch.id,
        source_artifact_id=artifact.id,
        structure_snapshot_id=structure.id,
        mapping_decision_id=decision.id,
        mapping_profile_usage_id=usage.id,
        source_artifact_generation=artifact.generation,
        mapping_decision_digest_sha256=decision.mapping_digest_sha256,
        profile_usage_mapping_digest_sha256=usage.mapping_digest_sha256,
        line_manifest=line_manifest,
        line_manifest_digest_sha256=service._sha256_hex(service._canonical_json(line_manifest)),
        finalized_by_user_id=actor.id,
        idempotency_key=f"preliminary-analysis-{suffix}",
        request_digest_sha256="e" * 64,
    )
    result_db.add(snapshot)
    result_db.commit()

    fake_storage = FakeObjectStorage()
    fake_storage._objects[artifact.storage_object_key] = source_data
    fake_storage._content_types[artifact.storage_object_key] = artifact.content_type
    set_object_storage_override(fake_storage)

    return {
        "org": org,
        "actor": actor,
        "role": role,
        "customer": customer,
        "project": project,
        "batch": batch,
        "artifact": artifact,
        "structure": structure,
        "decision": decision,
        "usage": usage,
        "snapshot": snapshot,
        "source_data": source_data,
        "fake_storage": fake_storage,
    }


def _generate(result_db: Session, seeded: dict, **overrides):
    values = {
        "actor": seeded["actor"],
        "org_id": seeded["org"].id,
        "project_id": seeded["project"].id,
        "preliminary_analysis_snapshot_id": seeded["snapshot"].id,
        "expected_project_version": seeded["project"].row_version,
        "idempotency_key": "preliminary-result-generate-1",
        "confirmed": True,
        "correlation_id": "corr-preliminary-result-1",
    }
    values.update(overrides)
    return generate_preliminary_result_artifact(result_db, **values)


def _assert_error(exc: pytest.ExceptionInfo[HTTPException], status: int, code: str) -> None:
    assert exc.value.status_code == status
    assert exc.value.detail["error_code"] == code


def test_generate_persists_artifact_and_atomic_audit(result_db: Session) -> None:
    seeded = _seed(result_db)

    artifact = _generate(result_db, seeded)

    assert artifact.project_id == seeded["project"].id
    assert artifact.version == 1
    assert artifact.created_by_user_id == seeded["actor"].id
    assert artifact.idempotency_key == "preliminary-result-generate-1"
    assert artifact.request_digest_sha256 is not None
    assert artifact.source_snapshot_sha256 == service._snapshot_canonical_digest(seeded["snapshot"])
    assert artifact.storage_object_key == f"org/{seeded['org'].id}/project/{seeded['project'].id}/preliminary-results/{artifact.id}.xlsx"
    assert artifact.original_filename == f"ket-qua-so-bo-v1-{artifact.id}.xlsx"

    audits = (
        result_db.query(AuditEvent)
        .filter(AuditEvent.event_name == "PreliminaryResultArtifactGenerated")
        .all()
    )
    assert len(audits) == 1
    assert audits[0].entity_id == artifact.id
    assert audits[0].command_name == "GeneratePreliminaryResultArtifact"


def test_output_contains_static_half_up_amounts(result_db: Session) -> None:
    seeded = _seed(
        result_db,
        rows=[("Máy bơm 10HP", 3.0), ("Máy bơm 20HP", 2.5)],
    )

    artifact = _generate(result_db, seeded)

    output_data = seeded["fake_storage"]._objects[artifact.storage_object_key]
    wb = load_workbook(io.BytesIO(output_data))
    try:
        ws = wb["Assets"]
        assert ws.cell(row=1, column=4).value == "Đơn giá đề xuất"
        assert ws.cell(row=1, column=5).value == "Thành tiền"
        # Row 3: qty 3.0 * 1050000 = 3150000.00
        assert ws.cell(row=3, column=4).value == Decimal("1050000.00")
        assert ws.cell(row=3, column=5).value == Decimal("3150000.00")
        # Row 4: qty 2.5 * 1050000 = 2625000.00
        assert ws.cell(row=4, column=4).value == Decimal("1050000.00")
        assert ws.cell(row=4, column=5).value == Decimal("2625000.00")
    finally:
        wb.close()


def test_same_idempotency_key_and_request_replays_without_new_artifact(
    result_db: Session,
) -> None:
    seeded = _seed(result_db)
    first = _generate(result_db, seeded)

    replay = _generate(result_db, seeded)

    assert replay.id == first.id
    assert result_db.query(PreliminaryResultArtifact).count() == 1
    assert (
        result_db.query(AuditEvent)
        .filter(AuditEvent.event_name == "PreliminaryResultArtifactGenerated")
        .count()
        == 1
    )


def test_missing_permission_denies_without_artifact(result_db: Session) -> None:
    seeded = _seed(result_db, grant_permission=False)

    with pytest.raises(HTTPException) as exc:
        _generate(result_db, seeded)

    _assert_error(exc, 403, "preliminary_result_forbidden")
    assert result_db.query(PreliminaryResultArtifact).count() == 0


@pytest.mark.parametrize("binding_state", ["inactive", "revoked"])
def test_inactive_or_revoked_role_binding_does_not_grant_permission(
    result_db: Session, binding_state: str
) -> None:
    seeded = _seed(result_db)
    if binding_state == "inactive":
        seeded["actor"].roles[0].is_active = False
    else:
        seeded["actor"].roles[0].revoked_at = datetime.now(timezone.utc)
    result_db.commit()

    with pytest.raises(HTTPException) as exc:
        _generate(result_db, seeded)

    _assert_error(exc, 403, "preliminary_result_forbidden")
    assert result_db.query(PreliminaryResultArtifact).count() == 0


def test_permission_is_rechecked_before_idempotent_replay(result_db: Session) -> None:
    seeded = _seed(result_db)
    artifact = _generate(result_db, seeded)
    seeded["actor"].roles[0].revoked_at = datetime.now(timezone.utc)
    result_db.commit()

    with pytest.raises(HTTPException) as exc:
        _generate(result_db, seeded)

    _assert_error(exc, 403, "preliminary_result_forbidden")
    assert result_db.query(PreliminaryResultArtifact).one().id == artifact.id


def test_reused_idempotency_key_with_different_request_is_rejected(
    result_db: Session,
) -> None:
    seeded = _seed(result_db)
    _generate(result_db, seeded)

    second_actor = User(
        organization_id=seeded["org"].id,
        email=f"preliminary-result-peer-{uuid.uuid4().hex[:8]}@example.com",
        full_name="Preliminary Result Peer",
        status=UserStatus.ACTIVE,
    )
    result_db.add(second_actor)
    result_db.flush()
    result_db.add(UserRole(user_id=second_actor.id, role_id=seeded["role"].id, is_active=True))
    result_db.commit()

    with pytest.raises(HTTPException) as exc:
        _generate(result_db, seeded, actor=second_actor)

    _assert_error(exc, 409, "preliminary_result_idempotency_key_reused")
    assert result_db.query(PreliminaryResultArtifact).count() == 1


def test_second_generate_for_project_is_rejected(result_db: Session) -> None:
    seeded = _seed(result_db)
    _generate(result_db, seeded)

    with pytest.raises(HTTPException) as exc:
        _generate(result_db, seeded, idempotency_key="preliminary-result-generate-2")

    _assert_error(exc, 409, "preliminary_result_already_generated")
    assert result_db.query(PreliminaryResultArtifact).count() == 1


def test_legacy_keyless_artifact_blocks_generation(result_db: Session) -> None:
    seeded = _seed(result_db)
    legacy = PreliminaryResultArtifact(
        organization_id=seeded["org"].id,
        customer_id=seeded["customer"].id,
        project_id=seeded["project"].id,
        version=1,
        original_filename="legacy.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        file_size_bytes=1,
        content_checksum_sha256="a" * 64,
        storage_object_key=f"legacy/{seeded['project'].id}.xlsx",
        source_snapshot_sha256="b" * 64,
        lineage_manifest={"contract": "legacy"},
        created_by_user_id=seeded["actor"].id,
    )
    result_db.add(legacy)
    result_db.commit()

    with pytest.raises(HTTPException) as exc:
        _generate(result_db, seeded)

    _assert_error(exc, 409, "preliminary_result_already_generated")


@pytest.mark.parametrize(
    ("idempotency_key", "request_digest_sha256"),
    [
        ("unpaired-key", None),
        (None, "a" * 64),
    ],
)
def test_unpaired_idempotency_key_and_digest_is_rejected_by_orm(
    result_db: Session,
    idempotency_key: str | None,
    request_digest_sha256: str | None,
) -> None:
    """Pairing CHECK rejects key without digest and digest without key at the ORM level."""
    seeded = _seed(result_db)
    artifact = PreliminaryResultArtifact(
        organization_id=seeded["org"].id,
        customer_id=seeded["customer"].id,
        project_id=seeded["project"].id,
        version=1,
        original_filename="unpaired.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        file_size_bytes=1,
        content_checksum_sha256="a" * 64,
        storage_object_key=f"unpaired/{seeded['project'].id}.xlsx",
        source_snapshot_sha256="b" * 64,
        lineage_manifest={"contract": "unpaired"},
        created_by_user_id=seeded["actor"].id,
        idempotency_key=idempotency_key,
        request_digest_sha256=request_digest_sha256,
    )
    result_db.add(artifact)
    with pytest.raises(IntegrityError):
        result_db.commit()
    result_db.rollback()


@pytest.mark.parametrize(
    ("overrides", "status", "code"),
    [
        ({"confirmed": False}, 400, "preliminary_result_confirmation_required"),
        ({"confirmed": 1}, 400, "preliminary_result_confirmation_required"),
        ({"confirmed": "yes"}, 400, "preliminary_result_confirmation_required"),
        ({"expected_project_version": 2}, 409, "project_version_conflict"),
    ],
)
def test_confirmation_and_version_fail_without_persistence(
    result_db: Session, overrides: dict, status: int, code: str
) -> None:
    seeded = _seed(result_db)

    with pytest.raises(HTTPException) as exc:
        _generate(result_db, seeded, **overrides)

    _assert_error(exc, status, code)
    assert result_db.query(PreliminaryResultArtifact).count() == 0
    assert result_db.query(AuditEvent).count() == 0


def test_cross_tenant_project_is_safe_not_found(result_db: Session) -> None:
    seeded = _seed(result_db, suffix="owner")
    other = _seed(result_db, suffix="other")

    with pytest.raises(HTTPException) as exc:
        _generate(
            result_db,
            seeded,
            actor=other["actor"],
            org_id=other["org"].id,
        )

    _assert_error(exc, 404, "project_not_found")
    assert result_db.query(PreliminaryResultArtifact).count() == 0


@pytest.mark.parametrize("inactive_target", ["actor", "organization"])
def test_inactive_actor_or_organization_is_forbidden(
    result_db: Session, inactive_target: str
) -> None:
    seeded = _seed(result_db)
    if inactive_target == "actor":
        seeded["actor"].status = UserStatus.INACTIVE
    else:
        seeded["org"].status = OrganizationStatus.INACTIVE
    result_db.commit()

    with pytest.raises(HTTPException) as exc:
        _generate(result_db, seeded)

    _assert_error(exc, 403, "preliminary_result_forbidden")
    assert result_db.query(PreliminaryResultArtifact).count() == 0


def test_source_artifact_not_available_is_rejected(result_db: Session) -> None:
    seeded = _seed(result_db, artifact_state=ImportSourceArtifactState.PENDING)

    with pytest.raises(HTTPException) as exc:
        _generate(result_db, seeded)

    _assert_error(exc, 409, "source_artifact_not_available")
    assert result_db.query(PreliminaryResultArtifact).count() == 0


def test_source_format_xls_is_rejected(result_db: Session) -> None:
    seeded = _seed(result_db)
    seeded["artifact"].detected_format = "xls"
    result_db.commit()

    with pytest.raises(HTTPException) as exc:
        _generate(result_db, seeded)

    _assert_error(exc, 409, "preliminary_result_source_format_not_xlsx")
    assert result_db.query(PreliminaryResultArtifact).count() == 0


def test_source_checksum_mismatch_is_rejected(result_db: Session) -> None:
    seeded = _seed(result_db)
    seeded["usage"].source_checksum_sha256 = "0" * 64
    result_db.commit()

    with pytest.raises(HTTPException) as exc:
        _generate(result_db, seeded)

    _assert_error(exc, 409, "preliminary_result_source_usage_checksum_mismatch")
    assert result_db.query(PreliminaryResultArtifact).count() == 0


def test_mapping_profile_usage_decision_mismatch_is_rejected(result_db: Session) -> None:
    seeded = _seed(result_db)
    other_decision = (
        result_db.query(ColumnMappingDecision)
        .filter(
            ColumnMappingDecision.project_id == seeded["project"].id,
            ColumnMappingDecision.id != seeded["decision"].id,
        )
        .one()
    )
    seeded["usage"].confirmation_decision_id = other_decision.id
    result_db.commit()

    with pytest.raises(HTTPException) as exc:
        _generate(result_db, seeded)

    _assert_error(exc, 409, "mapping_profile_usage_decision_mismatch")
    assert result_db.query(PreliminaryResultArtifact).count() == 0


def test_v1_shaped_snapshot_is_rejected_as_locator_missing(result_db: Session) -> None:
    seeded = _seed(result_db)
    seeded["snapshot"].line_manifest = [
        {
            "identity": "Máy bơm",
            "accepted_price_basis": "internet_survey",
            "confirmed_reference_price": 1_000_000.0,
            "transport_percentage": 5.0,
            "proposed_unit_price": 1_050_000.0,
            "human_line_confirmed": True,
            "has_unresolved_blocking_line": False,
        }
    ]
    result_db.commit()

    with pytest.raises(HTTPException) as exc:
        _generate(result_db, seeded)

    _assert_error(exc, 409, "preliminary_result_source_locator_missing")
    assert result_db.query(PreliminaryResultArtifact).count() == 0


def test_empty_stored_manifest_is_rejected_without_writes(result_db: Session) -> None:
    seeded = _seed(result_db)
    seeded["snapshot"].line_manifest = []
    seeded["snapshot"].line_manifest_digest_sha256 = service._sha256_hex(
        service._canonical_json(seeded["snapshot"].line_manifest)
    )
    result_db.commit()

    with pytest.raises(HTTPException) as exc:
        _generate(result_db, seeded)

    _assert_error(exc, 409, "preliminary_result_line_manifest_empty")
    assert result_db.query(PreliminaryResultArtifact).count() == 0
    assert result_db.query(AuditEvent).count() == 0
    # Deterministic object key should never have been written.
    snapshot_digest = service._snapshot_canonical_digest(seeded["snapshot"])
    request_digest = service._request_digest(
        actor_id=seeded["actor"].id,
        project_id=seeded["project"].id,
        snapshot_id=seeded["snapshot"].id,
        snapshot_digest=snapshot_digest,
        expected_project_version=seeded["project"].row_version,
    )
    expected_key = f"org/{seeded['org'].id}/project/{seeded['project'].id}/preliminary-results/{service._artifact_id(org_id=seeded['org'].id, normalized_key='preliminary-result-generate-1', request_digest=request_digest)}.xlsx"
    assert expected_key not in seeded["fake_storage"]._objects


def test_line_locator_out_of_region_is_rejected(result_db: Session) -> None:
    seeded = _seed(result_db)
    seeded["snapshot"].line_manifest = [_line(source_row_number=99)]
    seeded["snapshot"].line_manifest_digest_sha256 = service._sha256_hex(
        service._canonical_json(seeded["snapshot"].line_manifest)
    )
    result_db.commit()

    with pytest.raises(HTTPException) as exc:
        _generate(result_db, seeded)

    _assert_error(exc, 409, "preliminary_result_line_locator_out_of_region")
    assert result_db.query(PreliminaryResultArtifact).count() == 0


def test_target_not_empty_is_rejected(result_db: Session) -> None:
    seeded = _seed(result_db)
    # Write data into target column D within the candidate rows.
    wb = load_workbook(io.BytesIO(seeded["source_data"]))
    ws = wb["Assets"]
    ws.cell(row=3, column=4, value="blocked")
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    new_data = output.read()
    seeded["fake_storage"]._objects[seeded["artifact"].storage_object_key] = new_data
    seeded["artifact"].checksum_sha256 = _sha256(new_data)
    seeded["usage"].source_checksum_sha256 = seeded["artifact"].checksum_sha256
    result_db.commit()
    wb.close()

    with pytest.raises(HTTPException) as exc:
        _generate(result_db, seeded)

    _assert_error(exc, 409, "preliminary_result_target_not_empty")
    assert result_db.query(PreliminaryResultArtifact).count() == 0


def test_deterministic_artifact_id_and_digest(result_db: Session) -> None:
    seeded = _seed(result_db)
    snapshot_digest = service._snapshot_canonical_digest(seeded["snapshot"])
    request_digest = service._request_digest(
        actor_id=seeded["actor"].id,
        project_id=seeded["project"].id,
        snapshot_id=seeded["snapshot"].id,
        snapshot_digest=snapshot_digest,
        expected_project_version=seeded["project"].row_version,
    )
    expected_id = service._artifact_id(
        org_id=seeded["org"].id,
        normalized_key="preliminary-result-generate-1",
        request_digest=request_digest,
    )

    artifact = _generate(result_db, seeded)

    assert artifact.id == expected_id
    assert artifact.request_digest_sha256 == request_digest
    assert artifact.source_snapshot_sha256 == snapshot_digest


def test_lineage_manifest_shape_is_normative(result_db: Session) -> None:
    seeded = _seed(result_db)

    artifact = _generate(result_db, seeded)

    lineage = artifact.lineage_manifest
    assert lineage["generation_contract"] == "preliminary-result-generate-v1"
    assert lineage["project_id"] == str(seeded["project"].id)
    assert lineage["source_workbook"]["artifact_id"] == str(seeded["artifact"].id)
    assert lineage["analysis_snapshot"]["canonical_digest_sha256"] == artifact.source_snapshot_sha256
    assert lineage["output_layout"]["price_header"] == "Đơn giá đề xuất"
    assert lineage["output_layout"]["amount_header"] == "Thành tiền"
    assert lineage["output_layout"]["header_merge"] is True
    assert lineage["quantity_column"]["source_column_index"] == 2


def test_audit_failure_rolls_back_artifact_and_object(result_db: Session, monkeypatch) -> None:
    seeded = _seed(result_db)

    def fail_audit(*args, **kwargs):
        raise RuntimeError("audit unavailable")

    monkeypatch.setattr(service, "log_audit_event", fail_audit)
    with pytest.raises(RuntimeError, match="audit unavailable"):
        _generate(result_db, seeded)

    assert result_db.query(PreliminaryResultArtifact).count() == 0
    assert result_db.query(AuditEvent).count() == 0
    # The deterministic object should have been best-effort deleted.
    snapshot_digest = service._snapshot_canonical_digest(seeded["snapshot"])
    request_digest = service._request_digest(
        actor_id=seeded["actor"].id,
        project_id=seeded["project"].id,
        snapshot_id=seeded["snapshot"].id,
        snapshot_digest=snapshot_digest,
        expected_project_version=seeded["project"].row_version,
    )
    expected_id = service._artifact_id(
        org_id=seeded["org"].id,
        normalized_key="preliminary-result-generate-1",
        request_digest=request_digest,
    )
    storage_key = f"org/{seeded['org'].id}/project/{seeded['project'].id}/preliminary-results/{expected_id}.xlsx"
    assert storage_key not in seeded["fake_storage"]._objects


def test_storage_collision_mismatch_returns_500(result_db: Session) -> None:
    seeded = _seed(result_db)
    snapshot_digest = service._snapshot_canonical_digest(seeded["snapshot"])
    request_digest = service._request_digest(
        actor_id=seeded["actor"].id,
        project_id=seeded["project"].id,
        snapshot_id=seeded["snapshot"].id,
        snapshot_digest=snapshot_digest,
        expected_project_version=seeded["project"].row_version,
    )
    expected_id = service._artifact_id(
        org_id=seeded["org"].id,
        normalized_key="preliminary-result-generate-1",
        request_digest=request_digest,
    )
    storage_key = f"org/{seeded['org'].id}/project/{seeded['project'].id}/preliminary-results/{expected_id}.xlsx"
    # Pre-seed a different object at the deterministic key.
    seeded["fake_storage"]._objects[storage_key] = b"foreign-collision"

    with pytest.raises(HTTPException) as exc:
        _generate(result_db, seeded)

    _assert_error(exc, 500, "preliminary_result_storage_failure")
    assert result_db.query(PreliminaryResultArtifact).count() == 0
