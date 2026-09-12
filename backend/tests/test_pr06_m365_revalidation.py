"""Focused service and parser proofs for PR-06 OneDrive Personal revalidation."""

from __future__ import annotations

import copy
import hashlib
import io
import uuid
import zipfile
from datetime import datetime, timezone

import httpx
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi import HTTPException
from sqlalchemy.orm import Session

from app.modules.document_workspace.application.document_revision_service import (
    create_document_with_first_revision,
)
from app.modules.document_workspace.models import DocumentRevisionCurrentHead
from app.modules.m365_integration.application.bind_document_service import (
    bind_document_revision,
)
from app.modules.m365_integration.application.revalidation_service import (
    _definition_set_from_template_manifest,
    get_revalidation_readiness,
    revalidate_document,
    seal_existing_binding_baseline,
)
from app.modules.m365_integration.domain.managed_regions import (
    ManagedRegionDefinition,
    ManagedRegionDefinitionSet,
    ManagedRegionIntegrityError,
    fingerprint_docx,
)
from app.modules.m365_integration.infrastructure.graph_adapter import (
    MAX_DOWNLOAD_BYTES,
    MicrosoftGraphError,
    MicrosoftGraphGateway,
)
from app.modules.m365_integration.models import (
    M365ManagedContentBaseline,
    M365ManagedRegionBaseline,
    M365RevalidationObservation,
    M365RevisionBinding,
)
from app.modules.project_master_data.models import (
    AuditEvent,
    DocumentTemplate,
    EvidenceFile,
    GeneratedDocument,
    RenderJob,
    TemplateVersion,
)
from tests.test_pr05_m365_foundation import (
    FakeGraphGateway,
    FakeOAuthClient,
    TABLES,
    _connect,
    _seed,
    _vault,
)


PR06_TABLES = [
    EvidenceFile.__table__,
    DocumentTemplate.__table__,
    TemplateVersion.__table__,
    RenderJob.__table__,
    GeneratedDocument.__table__,
    M365ManagedContentBaseline.__table__,
    M365ManagedRegionBaseline.__table__,
    M365RevalidationObservation.__table__,
]
REGIONS = ManagedRegionDefinitionSet(
    authority_ref="template-version:test:v1",
    definitions=(
        ManagedRegionDefinition(
            region_key="appraised-value",
            locator="VALORA_APPRAISED_VALUE",
            semantic_type="text",
            normalization_contract="text-whitespace-v1",
        ),
    ),
)


@pytest.fixture
def pr06_db() -> Session:
    from sqlalchemy import create_engine, event
    from sqlalchemy.pool import StaticPool

    from app.db import Base

    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )

    @event.listens_for(engine, "connect")
    def set_sqlite_pragma(dbapi_connection, connection_record):
        del connection_record
        cursor = dbapi_connection.cursor()
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()

    Base.metadata.create_all(engine, tables=[*TABLES, *PR06_TABLES])
    session = Session(bind=engine)
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


def _docx(*, outside: str = "Narrative", managed: str = "1.000.000 VND") -> bytes:
    document = Document()
    document.add_paragraph(outside)
    sdt = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), "VALORA_APPRAISED_VALUE")
    properties.append(tag)
    content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = managed
    run.append(text)
    paragraph.append(run)
    content.append(paragraph)
    sdt.append(properties)
    sdt.append(content)
    document._body._element.append(sdt)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _rewrite_docx(
    content: bytes,
    *,
    replacements: dict[str, bytes] | None = None,
    extras: dict[str, bytes] | None = None,
) -> bytes:
    replacements = replacements or {}
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as source:
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                target.writestr(info.filename, replacements.get(info.filename, source.read(info)))
            for name, value in (extras or {}).items():
                target.writestr(name, value)
    return output.getvalue()


class RevalidationGraph(FakeGraphGateway):
    def __init__(self, content: bytes) -> None:
        super().__init__()
        self.content = content
        self.e_tag = '"etag-1"'
        self.c_tag = '"ctag-1"'
        self.content_calls = 0
        self.identity_override: str | None = None
        self.failure: Exception | None = None
        self.after_next_metadata_e_tag: str | None = None
        self.on_next_metadata = None

    def get_drive_item(self, *, access_token: str, drive_id: str, drive_item_id: str):
        if self.failure is not None:
            raise self.failure
        observed = super().get_drive_item(
            access_token=access_token,
            drive_id=drive_id,
            drive_item_id=drive_item_id,
        )
        result = observed.__class__(
            drive_id=observed.drive_id,
            drive_item_id=self.identity_override or observed.drive_item_id,
            graph_version_id=None,
            e_tag=self.e_tag,
            c_tag=self.c_tag,
            last_modified_at=datetime(2026, 9, 12, 8, 0, tzinfo=timezone.utc),
            size_bytes=len(self.content),
            name=observed.name,
            path=observed.path,
            web_url=observed.web_url,
        )
        if self.after_next_metadata_e_tag is not None:
            self.e_tag = self.after_next_metadata_e_tag
            self.after_next_metadata_e_tag = None
        if self.on_next_metadata is not None:
            callback = self.on_next_metadata
            self.on_next_metadata = None
            callback()
        return result

    def get_drive_item_content(
        self, *, access_token: str, drive_id: str, drive_item_id: str
    ) -> bytes:
        assert access_token == "ephemeral-refreshed-access-token"
        assert drive_id == self.drive.drive_id
        assert drive_item_id
        self.content_calls += 1
        return self.content


def _setup(pr06_db: Session) -> dict[str, object]:
    original = _docx()
    seeded = _seed(pr06_db, suffix=f"pr06-{uuid.uuid4().hex[:6]}")
    role = seeded["actor"].roles[0].role
    role.permissions = ["project:read", "project:update"]
    pr06_db.commit()
    graph = RevalidationGraph(original)
    connection, oauth, _ = _connect(
        pr06_db,
        seeded,
        oauth=FakeOAuthClient(),
        graph=graph,
    )
    revision = create_document_with_first_revision(
        pr06_db,
        actor=seeded["actor"],
        organization_id=seeded["organization"].id,
        project_id=seeded["project"].id,
        document_type="valuation_report",
        title="Báo cáo PR-06",
        data_snapshot_digest_sha256="a" * 64,
        content_checksum_sha256=hashlib.sha256(original).hexdigest(),
        idempotency_key=f"pr06-document-{uuid.uuid4()}",
    )
    binding = bind_document_revision(
        pr06_db,
        actor=seeded["actor"],
        organization_id=seeded["organization"].id,
        project_id=seeded["project"].id,
        document_id=revision.document_id,
        document_revision_id=revision.id,
        expected_document_revision=1,
        connection_id=connection.id,
        drive_item_id="pr06-item",
        idempotency_key=f"pr06-binding-{uuid.uuid4()}",
        oauth_client=oauth,
        graph_gateway=graph,
        credential_vault=_vault(pr06_db),
    )
    baseline = seal_existing_binding_baseline(
        pr06_db,
        actor=seeded["actor"],
        organization_id=seeded["organization"].id,
        project_id=seeded["project"].id,
        document_id=revision.document_id,
        expected_document_revision_id=revision.id,
        expected_document_revision=1,
        binding_id=binding.id,
        definition_set=REGIONS,
        idempotency_key=f"pr06-seal-{uuid.uuid4()}",
        oauth_client=oauth,
        graph_gateway=graph,
        credential_vault=_vault(pr06_db),
    )
    graph.content_calls = 0
    return {
        **seeded,
        "oauth": oauth,
        "graph": graph,
        "revision": revision,
        "binding": binding,
        "baseline": baseline,
        "original": original,
    }


def _revalidate(pr06_db: Session, context: dict[str, object], *, key: str):
    return revalidate_document(
        pr06_db,
        actor=context["actor"],
        organization_id=context["organization"].id,
        project_id=context["project"].id,
        document_id=context["revision"].document_id,
        expected_document_revision_id=context["revision"].id,
        expected_document_revision=1,
        trigger="explicit_refresh",
        idempotency_key=key,
        oauth_client=context["oauth"],
        graph_gateway=context["graph"],
        credential_vault=_vault(pr06_db),
    )


def test_docx_fingerprint_masks_managed_content_but_tracks_narrative() -> None:
    original = fingerprint_docx(_docx(), REGIONS)
    managed_change = fingerprint_docx(_docx(managed="2.000.000 VND"), REGIONS)
    outside_change = fingerprint_docx(_docx(outside="Narrative changed"), REGIONS)

    assert original.outside_managed_digest_sha256 == managed_change.outside_managed_digest_sha256
    assert original.regions[0].normalized_value_digest_sha256 != (
        managed_change.regions[0].normalized_value_digest_sha256
    )
    assert original.outside_managed_digest_sha256 != outside_change.outside_managed_digest_sha256
    assert original.regions[0].normalized_value_digest_sha256 == (
        outside_change.regions[0].normalized_value_digest_sha256
    )


def test_docx_fingerprint_rejects_unsafe_package_path() -> None:
    content = _rewrite_docx(_docx(), extras={"../escape.xml": b"<escape/>"})

    with pytest.raises(ManagedRegionIntegrityError, match="unsafe package path"):
        fingerprint_docx(content, REGIONS)


def test_docx_fingerprint_rejects_dtd_and_unsafe_external_relationship() -> None:
    original = _docx()
    with zipfile.ZipFile(io.BytesIO(original)) as archive:
        document_xml = archive.read("word/document.xml")
        relationships = archive.read("word/_rels/document.xml.rels")
    dtd_docx = _rewrite_docx(
        original,
        replacements={"word/document.xml": b"<!DOCTYPE document>" + document_xml},
    )
    unsafe_relationship = relationships.replace(
        b"</Relationships>",
        (
            b'<Relationship Id="unsafe" Type="urn:test" Target="file:///secret" '
            b'TargetMode="External"/></Relationships>'
        ),
    )
    relationship_docx = _rewrite_docx(
        original,
        replacements={"word/_rels/document.xml.rels": unsafe_relationship},
    )

    with pytest.raises(ManagedRegionIntegrityError, match="declarations are unsafe"):
        fingerprint_docx(dtd_docx, REGIONS)
    with pytest.raises(ManagedRegionIntegrityError, match="external relationship"):
        fingerprint_docx(relationship_docx, REGIONS)


@pytest.mark.parametrize("encoding", ["utf-16", "utf-16-le", "utf-16-be"])
def test_docx_fingerprint_rejects_utf16_entity_guard_bypass(encoding: str) -> None:
    original = _docx()
    entity_xml = (
        '<!DOCTYPE document [<!ENTITY payload "expanded">]>'
        '<document xmlns="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "&payload;</document>"
    ).encode(encoding)
    malicious = _rewrite_docx(
        original,
        replacements={"word/document.xml": entity_xml},
    )

    with pytest.raises(ManagedRegionIntegrityError, match="encoding is unsupported"):
        fingerprint_docx(malicious, REGIONS)


def test_docx_fingerprint_rejects_missing_and_duplicate_managed_region() -> None:
    missing = ManagedRegionDefinitionSet(
        authority_ref="template-version:test:missing",
        definitions=(
            ManagedRegionDefinition(
                region_key="missing",
                locator="VALORA_MISSING",
                semantic_type="text",
                normalization_contract="text-whitespace-v1",
            ),
        ),
    )
    duplicate_document = Document(io.BytesIO(_docx()))
    managed_element = duplicate_document._body._element.find(qn("w:sdt"))
    assert managed_element is not None
    duplicate_document._body._element.append(copy.deepcopy(managed_element))
    duplicate_output = io.BytesIO()
    duplicate_document.save(duplicate_output)

    with pytest.raises(ManagedRegionIntegrityError, match="missing"):
        fingerprint_docx(_docx(), missing)
    with pytest.raises(ManagedRegionIntegrityError, match="duplicated"):
        fingerprint_docx(duplicate_output.getvalue(), REGIONS)


def test_managed_region_definition_set_must_not_be_empty() -> None:
    with pytest.raises(ValueError, match="At least one"):
        ManagedRegionDefinitionSet(authority_ref="template:test", definitions=())


def test_template_manifest_resolver_accepts_only_explicit_versioned_shape() -> None:
    template_version_id = uuid.uuid4()
    definition_set = _definition_set_from_template_manifest(
        template_version_id=template_version_id,
        manifest={
            "managed_regions_contract": "valora-managed-regions-v1",
            "managed_regions": [
                {
                    "region_key": "appraised-value",
                    "locator": "VALORA_APPRAISED_VALUE",
                    "semantic_type": "text",
                    "normalization_contract": "text-whitespace-v1",
                }
            ],
        },
    )

    assert definition_set.authority_ref == (
        f"template-version:{template_version_id}:valora-managed-regions-v1"
    )
    assert definition_set.definitions[0].locator == "VALORA_APPRAISED_VALUE"

    with pytest.raises(ValueError, match="contract"):
        _definition_set_from_template_manifest(
            template_version_id=template_version_id,
            manifest={"managed_regions": []},
        )
    with pytest.raises(ValueError, match="entry"):
        _definition_set_from_template_manifest(
            template_version_id=template_version_id,
            manifest={
                "managed_regions_contract": "valora-managed-regions-v1",
                "managed_regions": [
                    {
                        "region_key": "forged",
                        "locator": "FORGED",
                        "semantic_type": "text",
                        "normalization_contract": "text-whitespace-v1",
                        "client_override": True,
                    }
                ],
            },
        )


def test_existing_binding_seal_is_immutable_and_digest_only(pr06_db: Session) -> None:
    context = _setup(pr06_db)
    baseline = context["baseline"]

    assert baseline.source_content_sha256 == hashlib.sha256(context["original"]).hexdigest()
    assert pr06_db.query(M365ManagedContentBaseline).count() == 1
    region = pr06_db.query(M365ManagedRegionBaseline).one()
    assert region.locator == "VALORA_APPRAISED_VALUE"
    assert "1.000.000" not in str(region.__dict__)
    assert context["graph"].content_calls == 0


def test_new_binding_seals_baseline_and_regions_in_same_commit(pr06_db: Session) -> None:
    original = _docx()
    seeded = _seed(pr06_db, suffix="new-atomic")
    seeded["actor"].roles[0].role.permissions = ["project:read", "project:update"]
    pr06_db.commit()
    graph = RevalidationGraph(original)
    connection, oauth, _ = _connect(pr06_db, seeded, graph=graph)
    revision = create_document_with_first_revision(
        pr06_db,
        actor=seeded["actor"],
        organization_id=seeded["organization"].id,
        project_id=seeded["project"].id,
        document_type="valuation_report",
        title="Atomic baseline",
        data_snapshot_digest_sha256="a" * 64,
        content_checksum_sha256=hashlib.sha256(original).hexdigest(),
        idempotency_key="new-atomic-document",
    )

    binding = bind_document_revision(
        pr06_db,
        actor=seeded["actor"],
        organization_id=seeded["organization"].id,
        project_id=seeded["project"].id,
        document_id=revision.document_id,
        document_revision_id=revision.id,
        expected_document_revision=1,
        connection_id=connection.id,
        drive_item_id="new-atomic-item",
        idempotency_key="new-atomic-binding",
        oauth_client=oauth,
        graph_gateway=graph,
        credential_vault=_vault(pr06_db),
        definition_set=REGIONS,
    )

    baseline = pr06_db.query(M365ManagedContentBaseline).one()
    assert baseline.binding_id == binding.id
    assert baseline.provenance_kind == "new_binding_atomic"
    assert baseline.source_content_sha256 == hashlib.sha256(original).hexdigest()
    assert pr06_db.query(M365ManagedRegionBaseline).count() == 1
    assert graph.item_calls == 2
    assert graph.content_calls == 1
    assert (
        pr06_db.query(AuditEvent).filter(AuditEvent.event_name == "M365_REVISION_BOUND").count()
        == 1
    )
    assert (
        pr06_db.query(AuditEvent)
        .filter(AuditEvent.event_name == "M365_MANAGED_CONTENT_BASELINE_SEALED")
        .count()
        == 1
    )


def test_new_binding_checksum_mismatch_rolls_back_binding_and_baseline(
    pr06_db: Session,
) -> None:
    original = _docx()
    seeded = _seed(pr06_db, suffix="new-mismatch")
    seeded["actor"].roles[0].role.permissions = ["project:read", "project:update"]
    pr06_db.commit()
    graph = RevalidationGraph(original)
    connection, oauth, _ = _connect(pr06_db, seeded, graph=graph)
    revision = create_document_with_first_revision(
        pr06_db,
        actor=seeded["actor"],
        organization_id=seeded["organization"].id,
        project_id=seeded["project"].id,
        document_type="valuation_report",
        title="Checksum mismatch",
        data_snapshot_digest_sha256="a" * 64,
        content_checksum_sha256="b" * 64,
        idempotency_key="new-mismatch-document",
    )

    with pytest.raises(HTTPException) as exc:
        bind_document_revision(
            pr06_db,
            actor=seeded["actor"],
            organization_id=seeded["organization"].id,
            project_id=seeded["project"].id,
            document_id=revision.document_id,
            document_revision_id=revision.id,
            expected_document_revision=1,
            connection_id=connection.id,
            drive_item_id="new-mismatch-item",
            idempotency_key="new-mismatch-binding",
            oauth_client=oauth,
            graph_gateway=graph,
            credential_vault=_vault(pr06_db),
            definition_set=REGIONS,
        )

    assert exc.value.status_code == 409
    assert exc.value.detail["error_code"] == "m365_content_checksum_mismatch"
    assert pr06_db.query(M365RevisionBinding).count() == 0
    assert pr06_db.query(M365ManagedContentBaseline).count() == 0


def test_binding_resolves_tenant_scope_before_provider_access(pr06_db: Session) -> None:
    seeded = _seed(pr06_db, suffix="bind-owner")
    foreign = _seed(pr06_db, suffix="bind-foreign")
    graph = RevalidationGraph(_docx())
    connection, oauth, _ = _connect(pr06_db, seeded, graph=graph)

    with pytest.raises(HTTPException) as exc:
        bind_document_revision(
            pr06_db,
            actor=seeded["actor"],
            organization_id=seeded["organization"].id,
            project_id=foreign["project"].id,
            document_id=uuid.uuid4(),
            document_revision_id=uuid.uuid4(),
            expected_document_revision=1,
            connection_id=connection.id,
            drive_item_id="foreign-item",
            idempotency_key="foreign-bind",
            oauth_client=oauth,
            graph_gateway=graph,
            credential_vault=_vault(pr06_db),
            definition_set=REGIONS,
        )

    assert exc.value.status_code == 404
    assert exc.value.detail["error_code"] == "project_not_found"
    assert oauth.refresh_calls == 0
    assert graph.item_calls == 0


def test_seal_rejects_changed_binding_metadata_without_writing(pr06_db: Session) -> None:
    original = _docx()
    seeded = _seed(pr06_db, suffix="seal-mismatch")
    seeded["actor"].roles[0].role.permissions = ["project:read", "project:update"]
    pr06_db.commit()
    graph = RevalidationGraph(original)
    connection, oauth, _ = _connect(pr06_db, seeded, graph=graph)
    revision = create_document_with_first_revision(
        pr06_db,
        actor=seeded["actor"],
        organization_id=seeded["organization"].id,
        project_id=seeded["project"].id,
        document_type="valuation_report",
        title="Mismatch",
        data_snapshot_digest_sha256="a" * 64,
        content_checksum_sha256=hashlib.sha256(original).hexdigest(),
        idempotency_key="seal-mismatch-document",
    )
    binding = bind_document_revision(
        pr06_db,
        actor=seeded["actor"],
        organization_id=seeded["organization"].id,
        project_id=seeded["project"].id,
        document_id=revision.document_id,
        document_revision_id=revision.id,
        expected_document_revision=1,
        connection_id=connection.id,
        drive_item_id="item",
        idempotency_key="seal-mismatch-binding",
        oauth_client=oauth,
        graph_gateway=graph,
        credential_vault=_vault(pr06_db),
    )
    graph.e_tag = '"etag-changed"'

    with pytest.raises(HTTPException) as exc:
        seal_existing_binding_baseline(
            pr06_db,
            actor=seeded["actor"],
            organization_id=seeded["organization"].id,
            project_id=seeded["project"].id,
            document_id=revision.document_id,
            expected_document_revision_id=revision.id,
            expected_document_revision=1,
            binding_id=binding.id,
            definition_set=REGIONS,
            idempotency_key="seal-mismatch",
            oauth_client=oauth,
            graph_gateway=graph,
            credential_vault=_vault(pr06_db),
        )

    assert exc.value.status_code == 409
    assert exc.value.detail["error_code"] == "revalidation_baseline_required"
    assert pr06_db.query(M365ManagedContentBaseline).count() == 0


@pytest.mark.parametrize(
    ("outside", "managed", "expected", "affected"),
    [
        ("Narrative changed", "1.000.000 VND", "external_change_outside_managed", []),
        ("Narrative", "2.000.000 VND", "external_change_in_managed", ["appraised-value"]),
        ("Both changed", "2.000.000 VND", "external_change_in_managed", ["appraised-value"]),
    ],
)
def test_revalidation_classifies_content_changes(
    pr06_db: Session,
    outside: str,
    managed: str,
    expected: str,
    affected: list[str],
) -> None:
    context = _setup(pr06_db)
    context["graph"].content = _docx(outside=outside, managed=managed)
    context["graph"].e_tag = '"etag-2"'
    context["graph"].c_tag = '"ctag-2"'

    observation = _revalidate(pr06_db, context, key=f"classify-{expected}")

    assert observation.classification == expected
    assert observation.affected_region_keys == affected
    assert context["graph"].content_calls == 1


def test_unchanged_etag_short_circuits_content_download(pr06_db: Session) -> None:
    context = _setup(pr06_db)

    observation = _revalidate(pr06_db, context, key="unchanged")

    assert observation.classification == "no_change"
    assert context["graph"].content_calls == 0


def test_rename_and_path_drift_keep_stable_identity_no_change(pr06_db: Session) -> None:
    context = _setup(pr06_db)
    context["graph"].item_name = "Renamed.docx"
    context["graph"].item_path = "/drive/root:/Moved/Renamed.docx"

    observation = _revalidate(pr06_db, context, key="rename-stable-id")

    assert observation.classification == "no_change"
    assert observation.observed_name == "Renamed.docx"
    assert observation.observed_path == "/drive/root:/Moved/Renamed.docx"
    assert context["graph"].content_calls == 0


def test_metadata_content_race_is_retryable_access_unavailable(pr06_db: Session) -> None:
    context = _setup(pr06_db)
    context["graph"].e_tag = '"etag-2"'
    context["graph"].after_next_metadata_e_tag = '"etag-3"'

    observation = _revalidate(pr06_db, context, key="metadata-content-race")

    assert observation.classification == "access_unavailable"
    assert observation.reason_category == "metadata_content_race"
    assert observation.retryable is True


def test_conflicting_idempotency_key_reuse_fails_before_second_provider_read(
    pr06_db: Session,
) -> None:
    context = _setup(pr06_db)
    _revalidate(pr06_db, context, key="reused-command")
    calls = context["graph"].item_calls

    with pytest.raises(HTTPException) as exc:
        revalidate_document(
            pr06_db,
            actor=context["actor"],
            organization_id=context["organization"].id,
            project_id=context["project"].id,
            document_id=context["revision"].document_id,
            expected_document_revision_id=context["revision"].id,
            expected_document_revision=1,
            trigger="reconnect",
            idempotency_key="reused-command",
            oauth_client=context["oauth"],
            graph_gateway=context["graph"],
            credential_vault=_vault(pr06_db),
        )

    assert exc.value.status_code == 409
    assert exc.value.detail["error_code"] == "idempotency_key_reused"
    assert context["graph"].item_calls == calls


@pytest.mark.parametrize("mutation", ["head", "binding"])
def test_lineage_change_during_provider_read_rejects_stale_observation(
    pr06_db: Session,
    mutation: str,
) -> None:
    context = _setup(pr06_db)
    context["graph"].e_tag = '"etag-2"'

    def mutate_lineage() -> None:
        if mutation == "head":
            head = pr06_db.query(DocumentRevisionCurrentHead).one()
            head.document_revision = 2
        else:
            binding = pr06_db.get(M365RevisionBinding, context["binding"].id)
            assert binding is not None
            binding.e_tag = '"binding-mutated"'
        pr06_db.commit()

    context["graph"].on_next_metadata = mutate_lineage

    with pytest.raises(HTTPException) as exc:
        _revalidate(pr06_db, context, key=f"lineage-race-{mutation}")

    assert exc.value.status_code == 409
    assert exc.value.detail["error_code"] in {
        "document_revision_conflict",
        "m365_revalidation_lineage_conflict",
    }
    assert pr06_db.query(M365RevalidationObservation).count() == 0


def test_missing_item_and_provider_failure_have_distinct_safe_results(
    pr06_db: Session,
) -> None:
    context = _setup(pr06_db)
    context["graph"].failure = MicrosoftGraphError(
        "Microsoft Graph item was not found.", category="not_found", retryable=False
    )
    missing = _revalidate(pr06_db, context, key="missing-item")
    assert missing.classification == "file_replaced_or_moved"
    assert missing.reason_category == "not_found"

    context["graph"].failure = MicrosoftGraphError(
        "Microsoft Graph is unavailable.", category="provider_unavailable", retryable=True
    )
    unavailable = _revalidate(pr06_db, context, key="provider-down")
    assert unavailable.classification == "access_unavailable"
    assert unavailable.retryable is True


def test_parser_failure_is_access_unavailable_and_never_no_change(pr06_db: Session) -> None:
    context = _setup(pr06_db)
    context["graph"].content = b"not-a-docx"
    context["graph"].e_tag = '"etag-bad"'

    observation = _revalidate(pr06_db, context, key="bad-docx")

    assert observation.classification == "access_unavailable"
    assert observation.reason_category == "content_integrity_unavailable"


@pytest.mark.parametrize(
    "field_name", ["parser_contract_version", "fingerprint_contract_version"]
)
def test_parser_contract_mismatch_fails_closed(
    pr06_db: Session, field_name: str
) -> None:
    context = _setup(pr06_db)
    setattr(context["baseline"], field_name, "retired-contract")
    pr06_db.commit()
    context["graph"].e_tag = '"etag-2"'

    observation = _revalidate(pr06_db, context, key=f"{field_name}-mismatch")

    assert observation.classification == "access_unavailable"
    assert observation.reason_category == "content_integrity_unavailable"
    assert context["graph"].content_calls == 0


def test_idempotent_replay_returns_one_observation_and_audit(pr06_db: Session) -> None:
    context = _setup(pr06_db)
    first = _revalidate(pr06_db, context, key="same-command")
    calls = context["graph"].item_calls
    second = _revalidate(pr06_db, context, key="same-command")

    assert first.id == second.id
    assert context["graph"].item_calls == calls
    assert pr06_db.query(M365RevalidationObservation).count() == 1
    assert (
        pr06_db.query(AuditEvent)
        .filter(AuditEvent.event_name == "M365_REVALIDATION_COMPLETED")
        .count()
        == 1
    )


def test_readiness_uses_only_current_lineage(pr06_db: Session) -> None:
    context = _setup(pr06_db)
    _revalidate(pr06_db, context, key="readiness")

    aggregate = get_revalidation_readiness(
        pr06_db,
        actor=context["actor"],
        organization_id=context["organization"].id,
        project_id=context["project"].id,
        document_id=context["revision"].document_id,
    )

    assert aggregate.baseline_eligible is True
    assert aggregate.classification == "no_change"
    assert aggregate.is_fresh is True
    assert aggregate.is_safe_for_freshness_required_action is True


@pytest.mark.parametrize(
    ("mode", "classification", "blocking_reason", "next_action", "retryable"),
    [
        (
            "managed",
            "external_change_in_managed",
            "managed_change_review_required",
            "review_managed_changes",
            False,
        ),
        (
            "missing",
            "file_replaced_or_moved",
            "binding_untrusted",
            "reconnect_or_rebind",
            False,
        ),
        (
            "unavailable",
            "access_unavailable",
            "access_unavailable",
            "retry_revalidation",
            True,
        ),
    ],
)
def test_blocking_results_are_not_fresh_or_safe(
    pr06_db: Session,
    mode: str,
    classification: str,
    blocking_reason: str,
    next_action: str,
    retryable: bool,
) -> None:
    context = _setup(pr06_db)
    if mode == "managed":
        context["graph"].content = _docx(managed="2.000.000 VND")
        context["graph"].e_tag = '"etag-2"'
    elif mode == "missing":
        context["graph"].failure = MicrosoftGraphError(
            "missing", category="not_found", retryable=False
        )
    else:
        context["graph"].failure = MicrosoftGraphError(
            "unavailable", category="provider_unavailable", retryable=True
        )
    _revalidate(pr06_db, context, key=f"readiness-{mode}")

    aggregate = get_revalidation_readiness(
        pr06_db,
        actor=context["actor"],
        organization_id=context["organization"].id,
        project_id=context["project"].id,
        document_id=context["revision"].document_id,
    )

    assert aggregate.classification == classification
    assert aggregate.is_fresh is False
    assert aggregate.is_safe_for_freshness_required_action is False
    assert aggregate.blocking_reason == blocking_reason
    assert aggregate.next_action == next_action
    assert aggregate.retryable is retryable


def test_foreign_tenant_fails_before_provider_access(pr06_db: Session) -> None:
    context = _setup(pr06_db)
    before = context["graph"].item_calls

    with pytest.raises(HTTPException) as exc:
        revalidate_document(
            pr06_db,
            actor=context["actor"],
            organization_id=uuid.uuid4(),
            project_id=context["project"].id,
            document_id=context["revision"].document_id,
            expected_document_revision_id=context["revision"].id,
            expected_document_revision=1,
            trigger="explicit_refresh",
            idempotency_key="foreign",
            oauth_client=context["oauth"],
            graph_gateway=context["graph"],
            credential_vault=_vault(pr06_db),
        )

    assert exc.value.status_code == 403
    assert context["graph"].item_calls == before


class _StreamResponse:
    def __init__(
        self,
        status_code: int,
        *,
        headers: dict[str, str] | None = None,
        chunks: tuple[bytes, ...] = (),
    ) -> None:
        self.status_code = status_code
        self.headers = headers or {}
        self._chunks = chunks

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc_value, traceback):
        del exc_type, exc_value, traceback
        return False

    def iter_bytes(self):
        yield from self._chunks


def test_graph_download_follows_https_redirect_without_forwarding_bearer(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls: list[tuple[str, dict[str, str]]] = []
    responses = iter(
        [
            _StreamResponse(
                302,
                headers={"location": "https://storage.example.test/file"},
            ),
            _StreamResponse(200, chunks=(b"part-1", b"-part-2")),
        ]
    )

    def fake_stream(method: str, url: str, **kwargs):
        assert method == "GET"
        calls.append((url, dict(kwargs["headers"])))
        return next(responses)

    monkeypatch.setattr(
        "app.modules.m365_integration.infrastructure.graph_adapter.httpx.stream",
        fake_stream,
    )

    content = MicrosoftGraphGateway().get_drive_item_content(
        access_token="provider-secret",
        drive_id="drive",
        drive_item_id="item",
    )

    assert content == b"part-1-part-2"
    assert calls[0][1] == {"Authorization": "Bearer provider-secret"}
    assert calls[1] == ("https://storage.example.test/file", {})


def test_graph_download_rejects_unsafe_redirect(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        "app.modules.m365_integration.infrastructure.graph_adapter.httpx.stream",
        lambda *args, **kwargs: _StreamResponse(
            302,
            headers={"location": "http://storage.example.test/file"},
        ),
    )

    with pytest.raises(MicrosoftGraphError) as exc:
        MicrosoftGraphGateway().get_drive_item_content(
            access_token="provider-secret",
            drive_id="drive",
            drive_item_id="item",
        )

    assert exc.value.category == "download_failed"
    assert exc.value.retryable is False


@pytest.mark.parametrize(
    "response",
    [
        _StreamResponse(200, headers={"content-length": str(MAX_DOWNLOAD_BYTES + 1)}),
        _StreamResponse(200, chunks=(b"x" * (MAX_DOWNLOAD_BYTES + 1),)),
    ],
)
def test_graph_download_enforces_declared_and_streamed_size_limits(
    monkeypatch: pytest.MonkeyPatch,
    response: _StreamResponse,
) -> None:
    monkeypatch.setattr(
        "app.modules.m365_integration.infrastructure.graph_adapter.httpx.stream",
        lambda *args, **kwargs: response,
    )

    with pytest.raises(MicrosoftGraphError) as exc:
        MicrosoftGraphGateway().get_drive_item_content(
            access_token="provider-secret",
            drive_id="drive",
            drive_item_id="item",
        )

    assert exc.value.category == "content_too_large"
    assert exc.value.retryable is False


def test_graph_download_rejects_partial_stream(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        "app.modules.m365_integration.infrastructure.graph_adapter.httpx.stream",
        lambda *args, **kwargs: _StreamResponse(
            200,
            headers={"content-length": "10"},
            chunks=(b"short",),
        ),
    )

    with pytest.raises(MicrosoftGraphError) as exc:
        MicrosoftGraphGateway().get_drive_item_content(
            access_token="provider-secret",
            drive_id="drive",
            drive_item_id="item",
        )

    assert exc.value.category == "download_failed"
    assert exc.value.retryable is True


class _MetadataResponse:
    def __init__(self, status_code: int) -> None:
        self.status_code = status_code

    def json(self):
        return {}


@pytest.mark.parametrize(
    ("status_code", "category", "retryable"),
    [(401, "access_denied", False), (429, "provider_request_failed", True)],
)
def test_graph_metadata_maps_expired_consent_and_throttling(
    monkeypatch: pytest.MonkeyPatch,
    status_code: int,
    category: str,
    retryable: bool,
) -> None:
    monkeypatch.setattr(
        "app.modules.m365_integration.infrastructure.graph_adapter.httpx.get",
        lambda *args, **kwargs: _MetadataResponse(status_code),
    )

    with pytest.raises(MicrosoftGraphError) as exc:
        MicrosoftGraphGateway().get_default_drive(access_token="provider-secret")

    assert exc.value.category == category
    assert exc.value.retryable is retryable


def test_graph_metadata_timeout_is_retryable_and_sanitized(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def timeout(*args, **kwargs):
        raise httpx.ReadTimeout("provider-sensitive-timeout")

    monkeypatch.setattr(
        "app.modules.m365_integration.infrastructure.graph_adapter.httpx.get",
        timeout,
    )

    with pytest.raises(MicrosoftGraphError) as exc:
        MicrosoftGraphGateway().get_default_drive(access_token="provider-secret")

    assert str(exc.value) == "Microsoft Graph is unavailable."
    assert exc.value.category == "provider_unavailable"
    assert exc.value.retryable is True
