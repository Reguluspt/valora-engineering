"""E1-E30 offline proofs for VALORA-ONEDRIVE-EXCHANGE-001."""
from __future__ import annotations

import asyncio
import hashlib
import io
from datetime import datetime, timezone
from types import SimpleNamespace

import pytest
import openpyxl
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi import HTTPException
from sqlalchemy import create_engine, event
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.db import Base
from app.modules.document_workspace.application.document_storage_service import (
    finalize_storage_revision,
)
from app.modules.document_workspace.domain.document_blob_store import (
    BlobReadStatus,
    InMemoryDocumentBlobStore,
)
from app.modules.document_workspace.models import (
    DocumentRevision,
    DocumentRevisionCurrentHead,
    DocumentStorageCandidate,
    DocumentStorageExecutionEvent,
    DocumentStorageExecutionIntent,
    DocumentStorageExecutionState,
    StorageObjectBinding,
)
from app.modules.excel_import.application import source_artifact_service
from app.modules.excel_import.infrastructure.object_storage import (
    FakeObjectStorage,
    set_object_storage_override,
)
from app.modules.excel_import.models import ImportSourceArtifact
from app.modules.m365_integration.application import exchange_import_service
from app.modules.m365_integration.application.connection_service import (
    get_connection_capabilities,
)
from app.modules.m365_integration.application.exchange_import_service import (
    create_docx_export,
    create_docx_working_copy,
    import_inbox_docx,
    import_inbox_xlsx,
    reimport_working_docx,
)
from app.modules.m365_integration.application.exchange_service import (
    execute_exchange_create,
    prepare_exchange_create,
    provision_exchange_namespace,
)
from app.modules.m365_integration.domain.exchange_fake import InMemoryM365GraphGateway
from app.modules.m365_integration.domain.graph_gateway import GraphMutationStatus
from app.modules.m365_integration.domain.managed_regions import (
    ManagedRegionDefinition,
    ManagedRegionDefinitionSet,
)
from app.modules.m365_integration.models import (
    M365ConnectionCapability,
    M365ConnectionGrantedScope,
    M365ExchangeArtifact,
    M365ExchangeOperation,
)
from app.modules.project_master_data.models import (
    AuditEvent,
    ImportBatchStatus,
    ProjectAssetImportBatch,
    ProjectAssetImportStagingRow,
)
from tests.test_pr05_m365_foundation import TABLES, _connect, _seed


REGIONS = ManagedRegionDefinitionSet(
    authority_ref="template-version:g8:v1",
    definitions=(
        ManagedRegionDefinition(
            region_key="appraised-value",
            locator="VALORA_APPRAISED_VALUE",
            semantic_type="text",
            normalization_contract="text-whitespace-v1",
        ),
    ),
)
RETENTION_ANCHOR = datetime(2026, 9, 20, tzinfo=timezone.utc)
RETENTION_UNTIL = datetime(2036, 9, 20, tzinfo=timezone.utc)
EXCHANGE_TABLES = [
    ProjectAssetImportBatch.__table__,
    ImportSourceArtifact.__table__,
    ProjectAssetImportStagingRow.__table__,
    DocumentStorageExecutionIntent.__table__,
    DocumentStorageCandidate.__table__,
    DocumentStorageExecutionState.__table__,
    DocumentStorageExecutionEvent.__table__,
    StorageObjectBinding.__table__,
    M365ExchangeArtifact.__table__,
    M365ExchangeOperation.__table__,
]


def _docx(*, outside: str = "Narrative", managed: str = "1.000.000 VND") -> bytes:
    document = Document()
    document.add_paragraph(outside)
    sdt = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    tag = OxmlElement("w:tag")
    tag.set(qn("w:val"), "VALORA_APPRAISED_VALUE")
    properties.append(tag)
    content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = managed
    run.append(text)
    paragraph.append(run)
    content.append(paragraph)
    sdt.append(properties)
    sdt.append(content)
    document._body._element.append(sdt)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _xlsx(*, asset_name: str = "Máy bơm") -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "PD-001"
    sheet.append(["STT", "Tên tài sản", "Đặc điểm", "ĐVT", "Số lượng", "Đơn giá"])
    sheet.append([1, asset_name, "Model X", "cái", 2, 100])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.fixture
def exchange_context():
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )

    @event.listens_for(engine, "connect")
    def set_sqlite_foreign_keys(dbapi_connection, connection_record):
        del connection_record
        cursor = dbapi_connection.cursor()
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()

    Base.metadata.create_all(engine, tables=[*TABLES, *EXCHANGE_TABLES])
    db = Session(bind=engine)
    seeded = _seed(db, suffix="g8")
    connection, _, _ = _connect(db, seeded)
    graph = InMemoryM365GraphGateway()
    connection.drive_id = graph.drive_id
    capability = db.query(M365ConnectionCapability).filter_by(
        organization_id=seeded["organization"].id,
        connection_id=connection.id,
        capability_code="APPFOLDER_WRITE_AVAILABLE",
    ).one()
    capability.available = True
    capability.evidence_scope = "files.readwrite.appfolder"
    capability.provenance = "oauth_grant"
    db.add(
        M365ConnectionGrantedScope(
            organization_id=seeded["organization"].id,
            user_id=seeded["actor"].id,
            connection_id=connection.id,
            normalized_scope="files.readwrite.appfolder",
            provenance="oauth_grant",
        )
    )
    db.commit()
    context = {
        "db": db,
        "engine": engine,
        "seeded": seeded,
        "connection": connection,
        "graph": graph,
        "blob": InMemoryDocumentBlobStore(),
        "excel_storage": FakeObjectStorage(),
    }
    set_object_storage_override(context["excel_storage"])
    try:
        yield context
    finally:
        set_object_storage_override(None)
        db.close()
        engine.dispose()


def _run(awaitable):
    return asyncio.run(awaitable)


def _seed_inbox(ctx, content: bytes, *, name: str = "Bao-cao.docx"):
    namespace = provision_exchange_namespace(
        graph_gateway=ctx["graph"], access_token="offline-token"
    )
    return ctx["graph"].seed_file(
        parent_item_id=namespace.inbox_item_id, name=name, content=content
    )


def _initial_docx(ctx, *, key: str = "g8-initial", content: bytes | None = None):
    content = content or _docx()
    item = _seed_inbox(ctx, content)
    seeded = ctx["seeded"]
    revision, binding, artifact = _run(
        import_inbox_docx(
            ctx["db"],
            actor=seeded["actor"],
            organization_id=seeded["organization"].id,
            project_id=seeded["project"].id,
            connection_id=ctx["connection"].id,
            drive_item_id=item.drive_item_id,
            document_type="valuation_report",
            title="Báo cáo G8",
            definition_set=REGIONS,
            idempotency_key=key,
            graph_gateway=ctx["graph"],
            access_token="offline-token",
            blob_store=ctx["blob"],
            storage_profile_id="g8-fake",
            container_name="g8-documents",
            retention_policy_code="official-document-10y",
            retention_anchor_at=RETENTION_ANCHOR,
            minimum_retain_until=RETENTION_UNTIL,
        )
    )
    return revision, binding, artifact


def _working(ctx, *, key: str = "g8-working"):
    revision, binding, _ = _initial_docx(ctx, key=f"{key}-initial")
    artifact = _run(
        create_docx_working_copy(
            ctx["db"],
            actor=ctx["seeded"]["actor"],
            organization_id=ctx["seeded"]["organization"].id,
            project_id=ctx["seeded"]["project"].id,
            connection_id=ctx["connection"].id,
            document_id=revision.document_id,
            destination_name="Bao-cao-working.docx",
            idempotency_key=key,
            graph_gateway=ctx["graph"],
            access_token="offline-token",
            blob_store=ctx["blob"],
        )
    )
    assert artifact is not None
    return revision, binding, artifact


def _head(ctx, document_id):
    return ctx["db"].query(DocumentRevisionCurrentHead).filter_by(
        organization_id=ctx["seeded"]["organization"].id,
        document_id=document_id,
    ).one()


def _assert_code(exc: pytest.ExceptionInfo[HTTPException], code: str):
    assert exc.value.detail["error_code"] == code


def _new_excel_batch(ctx, *, source_filename: str = "Nguon.xlsx"):
    batch = ProjectAssetImportBatch(
        organization_id=ctx["seeded"]["organization"].id,
        project_id=ctx["seeded"]["project"].id,
        source_filename=source_filename,
        created_by_user_id=ctx["seeded"]["actor"].id,
    )
    ctx["db"].add(batch)
    ctx["db"].commit()
    return batch


def _import_xlsx(ctx, *, content: bytes, batch, reimport: bool = False):
    item = _seed_inbox(ctx, content, name="Nguon.xlsx")
    artifact = import_inbox_xlsx(
        ctx["db"],
        actor=ctx["seeded"]["actor"],
        organization_id=ctx["seeded"]["organization"].id,
        project_id=ctx["seeded"]["project"].id,
        connection_id=ctx["connection"].id,
        batch_id=batch.id,
        drive_item_id=item.drive_item_id,
        graph_gateway=ctx["graph"],
        access_token="offline-token",
        request=SimpleNamespace(headers={}),
        reimport=reimport,
    )
    return artifact, item


def test_e1_explicit_docx_inbox_import(exchange_context):
    revision, _, artifact = _initial_docx(exchange_context)
    assert revision.document_revision == 1
    assert artifact.role == "inbox" and artifact.state == "IMPORTED"


def test_e2_xlsx_inbox_routes_real_source_and_staging_intake(exchange_context):
    ctx = exchange_context
    batch = _new_excel_batch(ctx)
    artifact, _ = _import_xlsx(ctx, content=_xlsx(), batch=batch)
    source = ctx["db"].get(ImportSourceArtifact, artifact.excel_source_artifact_id)
    ctx["db"].refresh(batch)
    staged = ctx["db"].query(ProjectAssetImportStagingRow).filter_by(
        import_batch_id=batch.id
    ).all()
    intake_audit = ctx["db"].query(AuditEvent).filter_by(
        entity_id=batch.id,
        event_name="ProjectAssetImportBatchUploaded",
    ).one()

    assert source is not None and source.state == "available"
    assert source.generation == 1
    assert batch.current_source_artifact_id == source.id
    assert batch.status == ImportBatchStatus.PARSED
    assert [row.proposed_asset_name for row in staged] == ["Máy bơm"]
    assert intake_audit.command_name == "ReplaceStagingRows"
    assert artifact.source_authority_type == "EXCEL_SOURCE_ARTIFACT"


def test_e8_excel_save_does_not_apply_or_replace_staging(exchange_context):
    ctx = exchange_context
    batch = _new_excel_batch(ctx)
    artifact, item = _import_xlsx(ctx, content=_xlsx(), batch=batch)
    source_id = artifact.excel_source_artifact_id
    staged_ids = [
        row.id
        for row in ctx["db"].query(ProjectAssetImportStagingRow).filter_by(
            import_batch_id=batch.id
        ).all()
    ]

    ctx["graph"].replace_file_content(
        drive_item_id=item.drive_item_id,
        content=_xlsx(asset_name="Máy phát đã lưu trong Excel"),
    )

    ctx["db"].refresh(batch)
    assert batch.current_source_artifact_id == source_id
    assert batch.status == ImportBatchStatus.PARSED
    assert [
        row.id
        for row in ctx["db"].query(ProjectAssetImportStagingRow).filter_by(
            import_batch_id=batch.id
        ).all()
    ] == staged_ids
    assert ctx["db"].query(AuditEvent).filter_by(
        entity_id=batch.id,
        event_name="ProjectAssetImportBatchApplied",
    ).count() == 0


def test_e3_import_persists_full_sha256(exchange_context):
    content = _docx(outside="E3")
    _, binding, artifact = _initial_docx(exchange_context, content=content)
    expected = hashlib.sha256(content).hexdigest()
    assert binding.content_sha256 == artifact.observed_sha256 == expected


def test_e4_import_persists_exact_byte_length(exchange_context):
    content = _docx(outside="E4")
    _, binding, artifact = _initial_docx(exchange_context, content=content)
    assert binding.byte_length == artifact.observed_byte_length == len(content)


def test_e5_duplicate_import_is_idempotent(exchange_context):
    ctx = exchange_context
    content = _docx(outside="E5")
    item = _seed_inbox(ctx, content)
    seeded = ctx["seeded"]
    kwargs = dict(
        actor=seeded["actor"],
        organization_id=seeded["organization"].id,
        project_id=seeded["project"].id,
        connection_id=ctx["connection"].id,
        drive_item_id=item.drive_item_id,
        document_type="valuation_report",
        title="E5",
        definition_set=REGIONS,
        idempotency_key="e5-key",
        graph_gateway=ctx["graph"],
        access_token="offline-token",
        blob_store=ctx["blob"],
        storage_profile_id="g8-fake",
        container_name="g8-documents",
        retention_policy_code="official-document-10y",
        retention_anchor_at=RETENTION_ANCHOR,
        minimum_retain_until=RETENTION_UNTIL,
    )
    first = _run(import_inbox_docx(ctx["db"], **kwargs))
    second = _run(import_inbox_docx(ctx["db"], **kwargs))
    assert second[0].id == first[0].id
    assert ctx["db"].query(DocumentRevision).count() == 1


def test_e6_working_create_leaves_authority_unchanged(exchange_context):
    revision, _, _ = _working(exchange_context)
    assert _head(exchange_context, revision.document_id).current_revision_id == revision.id
    assert exchange_context["db"].query(DocumentRevision).count() == 1


def test_e7_word_save_leaves_current_head_unchanged(exchange_context):
    revision, _, working = _working(exchange_context)
    exchange_context["graph"].replace_file_content(
        drive_item_id=working.drive_item_id, content=_docx(outside="Word save")
    )
    assert _head(exchange_context, revision.document_id).current_revision_id == revision.id


def test_e9_explicit_docx_reimport_creates_immutable_revision(exchange_context):
    ctx = exchange_context
    first, _, working = _working(ctx)
    ctx["graph"].replace_file_content(
        drive_item_id=working.drive_item_id, content=_docx(outside="Explicit reimport")
    )
    result = _run(
        reimport_working_docx(
            ctx["db"],
            actor=ctx["seeded"]["actor"],
            organization_id=ctx["seeded"]["organization"].id,
            project_id=ctx["seeded"]["project"].id,
            artifact_id=working.id,
            definition_set=REGIONS,
            idempotency_key="e9-reimport",
            graph_gateway=ctx["graph"],
            access_token="offline-token",
            blob_store=ctx["blob"],
            storage_profile_id="g8-fake",
            container_name="g8-documents",
            retention_policy_code="official-document-10y",
            retention_anchor_at=RETENTION_ANCHOR,
            minimum_retain_until=RETENTION_UNTIL,
        )
    )
    assert result.outcome == "REVISION_CREATED"
    assert result.revision.document_revision == 2
    assert _head(ctx, first.document_id).current_revision_id == result.revision.id


def test_e10_xlsx_reimport_creates_new_source_generation_and_staging(exchange_context):
    ctx = exchange_context
    batch = _new_excel_batch(ctx)
    first, item = _import_xlsx(ctx, content=_xlsx(), batch=batch)
    first_source_id = first.excel_source_artifact_id
    ctx["graph"].replace_file_content(
        drive_item_id=item.drive_item_id,
        content=_xlsx(asset_name="Máy phát thế hệ 2"),
    )

    second = import_inbox_xlsx(
        ctx["db"],
        actor=ctx["seeded"]["actor"],
        organization_id=ctx["seeded"]["organization"].id,
        project_id=ctx["seeded"]["project"].id,
        connection_id=ctx["connection"].id,
        batch_id=batch.id,
        drive_item_id=item.drive_item_id,
        graph_gateway=ctx["graph"],
        access_token="offline-token",
        request=SimpleNamespace(headers={}),
        reimport=True,
    )
    source = ctx["db"].get(ImportSourceArtifact, second.excel_source_artifact_id)
    ctx["db"].refresh(batch)
    staged = ctx["db"].query(ProjectAssetImportStagingRow).filter_by(
        import_batch_id=batch.id
    ).all()

    assert source is not None and source.generation == 2
    assert source.id != first_source_id
    assert batch.current_source_artifact_id == source.id
    assert [row.proposed_asset_name for row in staged] == ["Máy phát thế hệ 2"]


def test_e11_current_head_cas_loss_is_fail_closed(exchange_context, monkeypatch):
    ctx = exchange_context
    first, _, working = _working(ctx)
    ctx["graph"].replace_file_content(
        drive_item_id=working.drive_item_id, content=_docx(outside="CAS loss")
    )
    original_finalize = finalize_storage_revision

    def lose_cas(db, **kwargs):
        head = _head(ctx, first.document_id)
        competing = DocumentRevision(
            organization_id=ctx["seeded"]["organization"].id,
            project_id=ctx["seeded"]["project"].id,
            document_id=first.document_id,
            document_revision=2,
            data_snapshot_digest_sha256=first.data_snapshot_digest_sha256,
            content_checksum_sha256="c" * 64,
            idempotency_key="competing-e11",
            request_digest_sha256="d" * 64,
            created_by_user_id=ctx["seeded"]["actor"].id,
        )
        db.add(competing)
        db.flush()
        head.current_revision_id = competing.id
        head.document_revision = 2
        db.commit()
        return original_finalize(db, **kwargs)

    monkeypatch.setattr(exchange_import_service, "finalize_storage_revision", lose_cas)
    with pytest.raises(HTTPException) as exc:
        _run(
            reimport_working_docx(
                ctx["db"],
                actor=ctx["seeded"]["actor"],
                organization_id=ctx["seeded"]["organization"].id,
                project_id=ctx["seeded"]["project"].id,
                artifact_id=working.id,
                definition_set=REGIONS,
                idempotency_key="e11-reimport",
                graph_gateway=ctx["graph"],
                access_token="offline-token",
                blob_store=ctx["blob"],
                storage_profile_id="g8-fake",
                container_name="g8-documents",
                retention_policy_code="official-document-10y",
                retention_anchor_at=RETENTION_ANCHOR,
                minimum_retain_until=RETENTION_UNTIL,
            )
        )
    _assert_code(exc, "storage_head_superseded")
    assert ctx["db"].query(DocumentRevision).count() == 2


def test_e12_excel_current_source_conflict_fails_before_staging_or_linkage(
    exchange_context, monkeypatch
):
    ctx = exchange_context
    batch = _new_excel_batch(ctx)
    exchange_artifact, item = _import_xlsx(ctx, content=_xlsx(), batch=batch)
    initial_source_id = exchange_artifact.excel_source_artifact_id
    initial_staging_ids = [
        row.id
        for row in ctx["db"].query(ProjectAssetImportStagingRow).filter_by(
            import_batch_id=batch.id
        ).all()
    ]
    ctx["graph"].replace_file_content(
        drive_item_id=item.drive_item_id,
        content=_xlsx(asset_name="Nguồn cạnh tranh"),
    )
    monkeypatch.setattr(
        source_artifact_service,
        "_atomic_claim_current_pointer",
        lambda *args, **kwargs: False,
    )

    with pytest.raises(HTTPException) as exc:
        import_inbox_xlsx(
            ctx["db"],
            actor=ctx["seeded"]["actor"],
            organization_id=ctx["seeded"]["organization"].id,
            project_id=ctx["seeded"]["project"].id,
            connection_id=ctx["connection"].id,
            batch_id=batch.id,
            drive_item_id=item.drive_item_id,
            graph_gateway=ctx["graph"],
            access_token="offline-token",
            request=SimpleNamespace(headers={}),
            reimport=True,
        )

    _assert_code(exc, "exchange_excel_source_conflict")
    ctx["db"].refresh(batch)
    ctx["db"].refresh(exchange_artifact)
    assert batch.current_source_artifact_id == initial_source_id
    assert exchange_artifact.excel_source_artifact_id == initial_source_id
    assert [
        row.id
        for row in ctx["db"].query(ProjectAssetImportStagingRow).filter_by(
            import_batch_id=batch.id
        ).all()
    ] == initial_staging_ids
    assert ctx["db"].query(ImportSourceArtifact).filter_by(
        import_batch_id=batch.id,
        state="orphaned",
    ).count() == 1


def test_e13_export_leaves_authority_unchanged(exchange_context):
    ctx = exchange_context
    revision, _, _ = _initial_docx(ctx)
    exported = _run(
        create_docx_export(
            ctx["db"],
            actor=ctx["seeded"]["actor"],
            organization_id=ctx["seeded"]["organization"].id,
            project_id=ctx["seeded"]["project"].id,
            connection_id=ctx["connection"].id,
            document_id=revision.document_id,
            destination_name="Bao-cao-export.docx",
            idempotency_key="e13-export",
            graph_gateway=ctx["graph"],
            access_token="offline-token",
            blob_store=ctx["blob"],
        )
    )
    assert exported is not None and exported.role == "export"
    assert _head(ctx, revision.document_id).current_revision_id == revision.id


@pytest.mark.parametrize("mutation", ["rename", "move", "delete"])
def test_e14_e15_e16_working_provider_mutations_are_safe(exchange_context, mutation):
    ctx = exchange_context
    revision, _, working = _working(ctx)
    if mutation == "rename":
        ctx["graph"].rename_item(drive_item_id=working.drive_item_id, name="Renamed.docx")
    elif mutation == "move":
        ctx["graph"].move_item(
            drive_item_id=working.drive_item_id,
            parent_item_id=ctx["graph"].app_root_item_id,
        )
    else:
        ctx["graph"].delete_item(drive_item_id=working.drive_item_id)
    assert _head(ctx, revision.document_id).current_revision_id == revision.id


def test_e17_e18_provider_failure_isolates_authoritative_reads(exchange_context):
    ctx = exchange_context
    revision, binding, _ = _initial_docx(ctx)
    ctx["graph"].set_fault("before_send_unavailable")
    result = _run(
        create_docx_working_copy(
            ctx["db"],
            actor=ctx["seeded"]["actor"],
            organization_id=ctx["seeded"]["organization"].id,
            project_id=ctx["seeded"]["project"].id,
            connection_id=ctx["connection"].id,
            document_id=revision.document_id,
            destination_name="Unavailable.docx",
            idempotency_key="e17-unavailable",
            graph_gateway=ctx["graph"],
            access_token="offline-token",
            blob_store=ctx["blob"],
        )
    )
    assert result is None
    read = _run(
        ctx["blob"].read_verified(
            object_key=binding.object_key,
            expected_sha256=binding.content_sha256,
            expected_byte_length=binding.byte_length,
            max_bytes=binding.byte_length,
        )
    )
    assert read.status == BlobReadStatus.MATCH


def test_e19_ambiguous_create_reconciles_without_duplicate_mutation(exchange_context):
    ctx = exchange_context
    revision, _, _ = _initial_docx(ctx)
    ctx["graph"].set_fault("after_commit_unknown")
    artifact = _run(
        create_docx_working_copy(
            ctx["db"],
            actor=ctx["seeded"]["actor"],
            organization_id=ctx["seeded"]["organization"].id,
            project_id=ctx["seeded"]["project"].id,
            connection_id=ctx["connection"].id,
            document_id=revision.document_id,
            destination_name="Unknown.docx",
            idempotency_key="e19-unknown",
            graph_gateway=ctx["graph"],
            access_token="offline-token",
            blob_store=ctx["blob"],
        )
    )
    assert artifact is not None
    assert ctx["graph"].create_calls == 1


def test_e20_provider_event_never_directly_mutates_authority(exchange_context):
    revision, _, working = _working(exchange_context)
    exchange_context["graph"].replace_file_content(
        drive_item_id=working.drive_item_id, content=_docx(outside="event")
    )
    assert _head(exchange_context, revision.document_id).document_revision == 1


def test_e21_files_read_only_connection_cannot_write(exchange_context):
    ctx = exchange_context
    revision, _, _ = _initial_docx(ctx, key="e21-initial")
    cap = ctx["db"].query(M365ConnectionCapability).filter_by(
        organization_id=ctx["seeded"]["organization"].id,
        connection_id=ctx["connection"].id,
        capability_code="APPFOLDER_WRITE_AVAILABLE",
    ).one()
    cap.available = False
    cap.evidence_scope = None
    ctx["db"].commit()
    ensure_calls = ctx["graph"].ensure_calls
    with pytest.raises(HTTPException) as exc:
        _run(
            create_docx_working_copy(
                ctx["db"],
                actor=ctx["seeded"]["actor"],
                organization_id=ctx["seeded"]["organization"].id,
                project_id=ctx["seeded"]["project"].id,
                connection_id=ctx["connection"].id,
                document_id=revision.document_id,
                destination_name="Denied-high-level.docx",
                idempotency_key="e21-high-level",
                graph_gateway=ctx["graph"],
                access_token="offline-token",
                blob_store=ctx["blob"],
            )
        )
    _assert_code(exc, "exchange_reconsent_required")
    assert ctx["graph"].ensure_calls == ensure_calls
    with pytest.raises(HTTPException) as exc:
        prepare_exchange_create(
            ctx["db"],
            organization_id=ctx["seeded"]["organization"].id,
            project_id=ctx["seeded"]["project"].id,
            connection_id=ctx["connection"].id,
            idempotency_key="e21",
            operation_kind="CREATE_WORKING",
            target_role="working",
            media="docx",
            drive_id=ctx["connection"].drive_id,
            destination_parent_item_id="working",
            destination_name="Denied.docx",
            content=b"denied",
        )
    _assert_code(exc, "exchange_reconsent_required")


def test_e22_missing_appfolder_grant_fails_closed(exchange_context):
    ctx = exchange_context
    ctx["db"].query(M365ConnectionGrantedScope).filter_by(
        normalized_scope="files.readwrite.appfolder"
    ).delete()
    ctx["db"].commit()
    read_available, write_available = get_connection_capabilities(
        ctx["db"],
        organization_id=ctx["seeded"]["organization"].id,
        connection_id=ctx["connection"].id,
    )
    assert read_available is True
    assert write_available is False
    with pytest.raises(HTTPException) as exc:
        prepare_exchange_create(
            ctx["db"],
            organization_id=ctx["seeded"]["organization"].id,
            project_id=ctx["seeded"]["project"].id,
            connection_id=ctx["connection"].id,
            idempotency_key="e22",
            operation_kind="CREATE_WORKING",
            target_role="working",
            media="docx",
            drive_id=ctx["connection"].drive_id,
            destination_parent_item_id="working",
            destination_name="Denied.docx",
            content=b"denied",
        )
    _assert_code(exc, "exchange_reconsent_required")


def test_e23_appfolder_target_escape_rejected(exchange_context):
    ctx = exchange_context
    with pytest.raises(HTTPException) as exc:
        prepare_exchange_create(
            ctx["db"],
            organization_id=ctx["seeded"]["organization"].id,
            project_id=ctx["seeded"]["project"].id,
            connection_id=ctx["connection"].id,
            idempotency_key="e23",
            operation_kind="CREATE_WORKING",
            target_role="working",
            media="docx",
            drive_id="another-drive",
            destination_parent_item_id="escape",
            destination_name="Escape.docx",
            content=b"escape",
        )
    _assert_code(exc, "exchange_target_escape")

    namespace = provision_exchange_namespace(
        graph_gateway=ctx["graph"], access_token="offline-token"
    )
    escaped = prepare_exchange_create(
        ctx["db"],
        organization_id=ctx["seeded"]["organization"].id,
        project_id=ctx["seeded"]["project"].id,
        connection_id=ctx["connection"].id,
        idempotency_key="e23-same-drive",
        operation_kind="CREATE_WORKING",
        target_role="working",
        media="docx",
        drive_id=ctx["connection"].drive_id,
        destination_parent_item_id=namespace.app_root_item_id,
        destination_name="Escape-same-drive.docx",
        content=b"escape",
    )
    create_calls = ctx["graph"].create_calls
    with pytest.raises(HTTPException) as exc:
        execute_exchange_create(
            ctx["db"],
            organization_id=ctx["seeded"]["organization"].id,
            operation_id=escaped.id,
            content=b"escape",
            graph_gateway=ctx["graph"],
            access_token="offline-token",
            source_authority_type="PROVIDER_TRANSPORT",
        )
    _assert_code(exc, "exchange_target_escape")
    assert ctx["graph"].create_calls == create_calls


def test_e24_file_vs_folder_namespace_collision_fails(exchange_context):
    graph = exchange_context["graph"]
    graph.seed_file(parent_item_id=graph.app_root_item_id, name="VALORA", content=b"collision")
    with pytest.raises(HTTPException) as exc:
        provision_exchange_namespace(graph_gateway=graph, access_token="offline-token")
    _assert_code(exc, "exchange_namespace_collision")


def test_e25_same_key_different_digest_conflicts(exchange_context):
    ctx = exchange_context
    namespace = provision_exchange_namespace(
        graph_gateway=ctx["graph"], access_token="offline-token"
    )
    kwargs = dict(
        organization_id=ctx["seeded"]["organization"].id,
        project_id=ctx["seeded"]["project"].id,
        connection_id=ctx["connection"].id,
        idempotency_key="e25",
        operation_kind="CREATE_WORKING",
        target_role="working",
        media="docx",
        drive_id=ctx["connection"].drive_id,
        destination_parent_item_id=namespace.working_item_id,
        destination_name="E25.docx",
    )
    prepare_exchange_create(ctx["db"], content=b"first", **kwargs)
    with pytest.raises(HTTPException) as exc:
        prepare_exchange_create(ctx["db"], content=b"second", **kwargs)
    _assert_code(exc, "idempotency_key_reused")


def test_e26_provider_unknown_recovers_after_restart_without_create(exchange_context):
    ctx = exchange_context
    revision, binding, _ = _initial_docx(ctx)
    namespace = provision_exchange_namespace(
        graph_gateway=ctx["graph"], access_token="offline-token"
    )
    content = ctx["blob"].object_bytes(object_key=binding.object_key)
    operation = prepare_exchange_create(
        ctx["db"],
        organization_id=ctx["seeded"]["organization"].id,
        project_id=ctx["seeded"]["project"].id,
        connection_id=ctx["connection"].id,
        idempotency_key="e26",
        operation_kind="CREATE_WORKING",
        target_role="working",
        media="docx",
        drive_id=ctx["connection"].drive_id,
        destination_parent_item_id=namespace.working_item_id,
        destination_name="Restart.docx",
        content=content,
    )
    ctx["graph"].set_fault("after_commit_unknown")
    result = ctx["graph"].create_file(
        access_token="offline-token",
        drive_id=operation.drive_id,
        parent_item_id=operation.destination_parent_item_id,
        exact_name=operation.destination_name,
        content=content,
    )
    operation.state = "PROVIDER_UNKNOWN"
    operation.provider_request_id = result.provider_request_id
    ctx["db"].commit()
    artifact = execute_exchange_create(
        ctx["db"],
        organization_id=ctx["seeded"]["organization"].id,
        operation_id=operation.id,
        content=content,
        graph_gateway=ctx["graph"],
        access_token="offline-token",
        source_authority_type="DOCUMENT_REVISION",
        document_id=revision.document_id,
        document_revision_id=revision.id,
    )
    assert artifact is not None
    assert ctx["graph"].create_calls == 1


def test_created_item_verification_failure_persists_action_required(exchange_context):
    class CorruptAfterCreateGateway(InMemoryM365GraphGateway):
        def create_file(self, **kwargs):
            result = super().create_file(**kwargs)
            if result.status == GraphMutationStatus.CREATED and result.item is not None:
                self.replace_file_content(
                    drive_item_id=result.item.drive_item_id,
                    content=b"provider-corruption",
                )
            return result

    ctx = exchange_context
    graph = CorruptAfterCreateGateway(drive_id=ctx["connection"].drive_id)
    namespace = provision_exchange_namespace(
        graph_gateway=graph, access_token="offline-token"
    )
    operation = prepare_exchange_create(
        ctx["db"],
        organization_id=ctx["seeded"]["organization"].id,
        project_id=ctx["seeded"]["project"].id,
        connection_id=ctx["connection"].id,
        idempotency_key="created-corrupt",
        operation_kind="CREATE_WORKING",
        target_role="working",
        media="docx",
        drive_id=ctx["connection"].drive_id,
        destination_parent_item_id=namespace.working_item_id,
        destination_name="Corrupt.docx",
        content=b"expected-content",
    )

    with pytest.raises(HTTPException) as exc:
        execute_exchange_create(
            ctx["db"],
            organization_id=ctx["seeded"]["organization"].id,
            operation_id=operation.id,
            content=b"expected-content",
            graph_gateway=graph,
            access_token="offline-token",
            source_authority_type="PROVIDER_TRANSPORT",
        )

    _assert_code(exc, "exchange_provider_mismatch")
    ctx["db"].refresh(operation)
    assert operation.state == "FAILED_ACTION_REQUIRED"
    assert operation.failure_code == "provider_ambiguous_or_mismatch"


def test_e27_docx_initial_import_always_has_binding(exchange_context):
    revision, binding, _ = _initial_docx(exchange_context)
    assert binding.document_revision_id == revision.id
    assert exchange_context["db"].query(StorageObjectBinding).count() == 1


def test_e28_xlsx_real_import_and_reimport_never_construct_document_revision(
    exchange_context,
):
    ctx = exchange_context
    batch = _new_excel_batch(ctx)
    _, item = _import_xlsx(ctx, content=_xlsx(), batch=batch)
    ctx["graph"].replace_file_content(
        drive_item_id=item.drive_item_id,
        content=_xlsx(asset_name="Máy phát không phải DOCX"),
    )
    import_inbox_xlsx(
        ctx["db"],
        actor=ctx["seeded"]["actor"],
        organization_id=ctx["seeded"]["organization"].id,
        project_id=ctx["seeded"]["project"].id,
        connection_id=ctx["connection"].id,
        batch_id=batch.id,
        drive_item_id=item.drive_item_id,
        graph_gateway=ctx["graph"],
        access_token="offline-token",
        request=SimpleNamespace(headers={}),
        reimport=True,
    )
    assert ctx["db"].query(DocumentRevision).count() == 0


def test_e29_onedrive_xlsx_flow_stops_at_staging_until_explicit_apply(
    exchange_context,
):
    ctx = exchange_context
    batch = _new_excel_batch(ctx)
    _import_xlsx(ctx, content=_xlsx(), batch=batch)
    ctx["db"].refresh(batch)

    assert batch.status == ImportBatchStatus.PARSED
    assert ctx["db"].query(ProjectAssetImportStagingRow).filter_by(
        import_batch_id=batch.id
    ).count() == 1
    assert ctx["db"].query(AuditEvent).filter_by(
        entity_id=batch.id,
        event_name="ProjectAssetImportBatchApplied",
    ).count() == 0


def test_e30_exchange_artifact_deletion_never_deletes_authoritative_blob(exchange_context):
    ctx = exchange_context
    _, binding, artifact = _initial_docx(ctx)
    object_key = binding.object_key
    ctx["db"].delete(artifact)
    ctx["db"].commit()
    assert ctx["blob"].object_bytes(object_key=object_key) is not None
